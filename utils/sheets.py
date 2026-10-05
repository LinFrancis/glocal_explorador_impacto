# -*- coding: utf-8 -*-
"""Almacén en Google Sheets: la base de datos es una hoja de cálculo de Google, no un archivo.

Funciona igual que `LocalStorage` (misma interfaz) y reutiliza toda la lógica de `Libro`/`repo`: guarda una
copia en memoria del contenido de la hoja, entrega a cada transacción un libro de trabajo de openpyxl y, al
terminar, sube solo las filas que cambiaron. Las celdas se guardan como TEXTO canónico (fechas ISO, «True»/«False»,
números con punto), así la hoja se lee y se edita a mano sin sorpresas de formato.

Configuración (`.streamlit/secrets.toml`, o los secretos de Streamlit Cloud):

    [gsheets]
    spreadsheet = "https://docs.google.com/spreadsheets/d/<ID>/edit"      # o solo el ID
    type = "service_account"
    project_id = "..."
    private_key_id = "..."
    private_key = "-----BEGIN PRIVATE KEY-----\\n...\\n-----END PRIVATE KEY-----\\n"
    client_email = "...@....iam.gserviceaccount.com"
    client_id = "..."
    token_uri = "https://oauth2.googleapis.com/token"

La hoja debe estar compartida (como Editor) con el `client_email`. La primera vez, si está vacía, se llena con el
libro local `data/catalogo_gestion.xlsx` (si existe) o con el Excel original.

El acceso a la API está detrás de `ClienteHojas`, así las pruebas usan un cliente en memoria (`ClienteMemoria`).
"""
from __future__ import annotations

import io
import os
import re
import threading
import time
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

from utils import schema as S
from utils import storage
from utils import variables as V
from utils.storage import AlmacenBase, ErrorAlmacen, Libro, Respaldo, _migrar
from utils.validation import a_texto, de_texto

MAX_CELDA = 32000                       # límite de openpyxl/Excel por celda (32 767); Google Sheets admite hasta 50 000
PREFIJO_RESPALDO = "RESPALDO_"
HOJAS_POR_DEFECTO = {"hoja 1", "hoja1", "sheet1", "sheet 1"}
MAX_FILAS_PARCIAL = 400                 # si cambian más filas que esto, se reescribe la hoja completa
BLOQUE_CARACTERES = 1_500_000           # tamaño máximo (aprox.) de cada petición de escritura

Matriz = list[list[str]]


# ============================================================================ configuración
SECCIONES = ("gsheets", ("connections", "gsheets"))
CLAVES_SPREADSHEET = ("spreadsheet", "spreadsheet_url", "spreadsheet_id", "url")


def _seccion_de_secretos():
    import streamlit as st
    try:
        secretos = st.secrets
        for sec in SECCIONES:
            nodo = secretos
            for parte in ((sec,) if isinstance(sec, str) else sec):
                nodo = nodo[parte]
            return dict(nodo)
    except Exception:  # noqa: BLE001  (sin secrets.toml o sin esa sección)
        return None
    return None


def configuracion() -> dict | None:
    """{'spreadsheet': ..., 'info': {cuenta de servicio}} si los secretos configuran Google Sheets; None si no.

    `IMPACTO_ALMACEN=local` fuerza el archivo local (pruebas). Si la sección existe pero está incompleta se avisa
    con un error claro, en vez de caer en silencio a un archivo local (en la nube se perdería al reiniciar).
    """
    if os.environ.get("IMPACTO_ALMACEN") == "local":
        return None
    sec = _seccion_de_secretos()
    if not sec:
        return None
    hoja = next((str(sec[k]).strip() for k in CLAVES_SPREADSHEET if sec.get(k)), "")
    info = {k: v for k, v in sec.items() if k not in CLAVES_SPREADSHEET}
    faltan = [k for k in ("client_email", "private_key") if not info.get(k)]
    if faltan or not hoja:
        partes = ([f"falta la dirección de la hoja (`spreadsheet`)"] if not hoja else []) + \
                 ([f"faltan los datos de la cuenta de servicio ({', '.join(faltan)})"] if faltan else [])
        raise ErrorAlmacen("La sección [gsheets] de los secretos está incompleta: " + " y ".join(partes) + ".")
    if "\\n" in str(info["private_key"]) and "\n" not in str(info["private_key"]):
        info["private_key"] = str(info["private_key"]).replace("\\n", "\n")
    info.setdefault("type", "service_account")
    info.setdefault("token_uri", "https://oauth2.googleapis.com/token")
    return {"spreadsheet": hoja, "info": info}


def crear(cfg: dict, semilla: Path | None, inicial: Path | None) -> "SheetsStorage":
    return SheetsStorage(ClienteGspread(cfg["info"], cfg["spreadsheet"]), semilla, inicial)


# ============================================================================ clientes
class ClienteHojas:
    """Lo mínimo que el almacén necesita de la hoja de cálculo (todo en texto)."""

    url = ""

    def marca(self) -> str:
        """Cambia cada vez que alguien modifica la hoja."""
        raise NotImplementedError

    def leer(self) -> dict[str, list[list[str]]]:
        """Todas las pestañas (sin las de respaldo): {título: filas de texto}."""
        raise NotImplementedError

    def titulos(self) -> list[str]:
        raise NotImplementedError

    def reemplazar(self, hoja: str, filas: Matriz) -> None:
        """Crea la pestaña si falta y deja exactamente `filas`."""
        raise NotImplementedError

    def escribir_filas(self, hoja: str, filas: dict[int, list[str]], ancho: int) -> None:
        """Sobrescribe filas sueltas ({n.º de fila desde 1: celdas})."""
        raise NotImplementedError

    def borrar_hoja(self, hoja: str) -> None:
        raise NotImplementedError


class ClienteMemoria(ClienteHojas):
    """Hoja de cálculo simulada en memoria (para pruebas): mismas operaciones, sin red."""

    url = "https://docs.google.com/spreadsheets/d/memoria/edit"

    def __init__(self, hojas: dict[str, Matriz] | None = None):
        self.hojas: dict[str, Matriz] = {k: [list(f) for f in v] for k, v in (hojas or {}).items()}
        self.version = 0
        self.llamadas = {"marca": 0, "leer": 0, "reemplazar": 0, "escribir_filas": 0, "borrar_hoja": 0}
        self.filas_escritas = 0

    def marca(self) -> str:
        self.llamadas["marca"] += 1
        return str(self.version)

    def leer(self):
        self.llamadas["leer"] += 1
        return {k: [list(f) for f in v] for k, v in self.hojas.items() if not k.startswith(PREFIJO_RESPALDO)}

    def titulos(self):
        return list(self.hojas)

    def reemplazar(self, hoja, filas):
        self.llamadas["reemplazar"] += 1
        self.hojas[hoja] = [list(f) for f in filas]
        self.filas_escritas += len(filas)
        self.version += 1

    def escribir_filas(self, hoja, filas, ancho):
        self.llamadas["escribir_filas"] += 1
        m = self.hojas.setdefault(hoja, [])
        for n, celdas in sorted(filas.items()):
            while len(m) < n:
                m.append([])
            m[n - 1] = list(celdas)
        self.filas_escritas += len(filas)
        self.version += 1

    def borrar_hoja(self, hoja):
        self.llamadas["borrar_hoja"] += 1
        self.hojas.pop(hoja, None)
        self.version += 1

    def editar_a_mano(self, hoja: str, fila: int, columna: int, texto: str) -> None:
        """Simula que alguien cambia una celda directamente en Google Sheets."""
        m = self.hojas[hoja]
        while len(m[fila - 1]) < columna:
            m[fila - 1].append("")
        m[fila - 1][columna - 1] = texto
        self.version += 1


def _reintentar(fn, *args, **kwargs):
    """Reintenta con espera creciente ante límites de cuota (429) y fallos pasajeros del servicio (5xx)."""
    import gspread
    espera = 2.0
    for intento in range(5):
        try:
            return fn(*args, **kwargs)
        except gspread.exceptions.APIError as e:
            codigo = getattr(getattr(e, "response", None), "status_code", 0)
            if codigo in (429, 500, 502, 503, 504) and intento < 4:
                time.sleep(espera)
                espera *= 2
                continue
            raise


class ClienteGspread(ClienteHojas):
    """Cliente real, sobre la API de Google Sheets (biblioteca gspread) con una cuenta de servicio."""

    def __init__(self, info: dict, spreadsheet: str):
        import gspread
        self._gspread = gspread
        self.correo = info.get("client_email", "")
        try:
            gc = gspread.service_account_from_dict(info)
            self.ss = _reintentar(gc.open_by_url if spreadsheet.startswith("http") else gc.open_by_key, spreadsheet)
        except Exception as e:  # noqa: BLE001
            raise ErrorAlmacen(self._explicar(e)) from e
        self.url = f"https://docs.google.com/spreadsheets/d/{self.ss.id}/edit"
        self._ws: dict[str, object] = {}

    def _explicar(self, e: Exception) -> str:
        g = self._gspread.exceptions
        if isinstance(e, g.SpreadsheetNotFound):
            return ("No se encuentra la hoja de Google. Revisa la dirección en los secretos y comparte la hoja, como Editor, "
                    f"con {self.correo or 'la cuenta de servicio'}.")
        if isinstance(e, g.APIError):
            codigo = getattr(getattr(e, "response", None), "status_code", 0)
            if codigo in (401, 403):
                return ("Google no deja acceder a la hoja. Comparte la hoja, como Editor, con "
                        f"{self.correo or 'la cuenta de servicio'} y revisa que la API de Google Sheets esté habilitada en su proyecto.")
            if codigo == 400 and "not supported" in str(e).lower():
                return ("Ese archivo de Drive es un Excel (.xlsx), no una hoja de cálculo de Google. Ábrelo en Drive y usa "
                        "Archivo → Guardar como hoja de cálculo de Google; luego pon la dirección de la hoja nueva en los secretos.")
            if codigo == 404:
                return "No se encuentra la hoja de Google: revisa su dirección en los secretos."
            if codigo == 429:
                return "Google limitó temporalmente las consultas. Espera un minuto e inténtalo de nuevo."
            return f"Google respondió con un error ({codigo}). Inténtalo de nuevo en un momento."
        if isinstance(e, (ValueError, KeyError)):
            return "Los datos de la cuenta de servicio en los secretos no son válidos (revisa private_key y client_email)."
        return f"No se pudo conectar con Google Sheets: {e}"

    def _envolver(self, fn, *args, **kwargs):
        try:
            return _reintentar(fn, *args, **kwargs)
        except ErrorAlmacen:
            raise
        except Exception as e:  # noqa: BLE001
            raise ErrorAlmacen(self._explicar(e)) from e

    def marca(self) -> str:
        try:
            return str(_reintentar(self.ss.get_lastUpdateTime))
        except Exception:  # noqa: BLE001  (sin la API de Drive: se vuelve a leer cada 30 segundos)
            return f"t{int(time.time() // 30)}"

    def leer(self):
        hojas = self._envolver(self.ss.worksheets)
        self._ws = {w.title: w for w in hojas}
        titulos = [w.title for w in hojas if not w.title.startswith(PREFIJO_RESPALDO)]
        if not titulos:
            return {}
        res = self._envolver(self.ss.values_batch_get, [f"'{t}'" for t in titulos])
        return {t: vr.get("values", []) for t, vr in zip(titulos, res.get("valueRanges", []))}

    def titulos(self):
        return [w.title for w in self._envolver(self.ss.worksheets)]

    def _hoja(self, titulo: str):
        if titulo not in self._ws:
            self._ws = {w.title: w for w in self._envolver(self.ss.worksheets)}
        return self._ws.get(titulo)

    @staticmethod
    def _bloques(filas: Matriz):
        """Reparte las filas en bloques de tamaño razonable para cada petición."""
        bloque, peso, inicio = [], 0, 0
        for i, f in enumerate(filas):
            p = sum(len(c) for c in f) + len(f)
            if bloque and peso + p > BLOQUE_CARACTERES:
                yield inicio, bloque
                bloque, peso, inicio = [], 0, i
            bloque.append(f)
            peso += p
        if bloque:
            yield inicio, bloque

    def reemplazar(self, hoja, filas):
        ancho = max((len(f) for f in filas), default=1)
        n_filas, n_cols = max(len(filas) + 20, 100), max(ancho, 10)
        ws = self._hoja(hoja)
        if ws is None:
            ws = self._envolver(self.ss.add_worksheet, title=hoja, rows=n_filas, cols=n_cols)
            self._ws[hoja] = ws
        else:
            self._envolver(ws.clear)
            self._envolver(ws.resize, rows=n_filas, cols=n_cols)
        for inicio, bloque in self._bloques(filas):
            self._envolver(ws.update, values=bloque, range_name=f"A{inicio + 1}", raw=True)

    def escribir_filas(self, hoja, filas, ancho):
        ws = self._hoja(hoja)
        if ws is None:
            raise ErrorAlmacen(f"No existe la pestaña «{hoja}» en la hoja de Google.")
        ultima = max(filas)
        if ws.row_count < ultima + 20 or ws.col_count < ancho:
            self._envolver(ws.resize, rows=max(ws.row_count, ultima + 100), cols=max(ws.col_count, ancho))
        # filas consecutivas -> un solo rango
        datos, actual = [], None
        for n in sorted(filas):
            if actual and n == actual[0] + len(actual[1]):
                actual[1].append(filas[n])
            else:
                actual = [n, [filas[n]]]
                datos.append(actual)
        peticion: list[dict] = []
        peso = 0
        for n, bloque in datos:
            p = sum(len(c) + 1 for f in bloque for c in f)
            if peticion and peso + p > BLOQUE_CARACTERES:
                self._envolver(self.ss.values_batch_update, {"valueInputOption": "RAW", "data": peticion})
                peticion, peso = [], 0
            peticion.append({"range": f"'{hoja}'!A{n}", "values": bloque})
            peso += p
        if peticion:
            self._envolver(self.ss.values_batch_update, {"valueInputOption": "RAW", "data": peticion})

    def borrar_hoja(self, hoja):
        ws = self._hoja(hoja)
        if ws is not None:
            self._envolver(self.ss.del_worksheet, ws)
            self._ws.pop(hoja, None)


# ============================================================================ conversión libro <-> matrices
def _tipo_columna(hoja: str, cab: str) -> str | None:
    c = S.CAMPO.get(cab) if hoja == S.HOJA_BASE else None
    return c.tipo if c else None


def matriz_de_hoja(ws, nombre: str) -> Matriz:
    """Hoja de openpyxl -> filas de texto canónico, sin columnas ni filas vacías sobrantes al final."""
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cab = [str(v).strip() if v is not None else "" for v in filas[0]]
    ancho = max((i + 1 for i, h in enumerate(cab) if h), default=0)
    if ancho == 0:
        return []
    tipos = [_tipo_columna(nombre, h) for h in cab[:ancho]]
    out: Matriz = [cab[:ancho]]
    for fila in filas[1:]:
        valores = list(fila[:ancho]) + [None] * (ancho - len(fila))
        out.append([a_texto(v, t)[:MAX_CELDA] for v, t in zip(valores, tipos)])
    while len(out) > 1 and not any(out[-1]):
        out.pop()
    return out


def normalizar(filas: list[list]) -> Matriz:
    """Lo que devuelve la API (filas de distinto largo) -> matriz rectangular de texto."""
    if not filas:
        return []
    ancho = max((i + 1 for i, h in enumerate(filas[0]) if str(h).strip()), default=0)
    if ancho == 0:
        return []
    out = [[str(c).strip() for c in (list(f) + [""] * ancho)[:ancho]] for f in filas]
    while len(out) > 1 and not any(out[-1]):
        out.pop()
    return out


def _valor_celda(texto: str, tipo: str | None):
    """Valor tipado de una celda de texto, solo si convertirlo no cambia lo escrito en la hoja
    («2026-07-14T14:54:37.77» o «sí» se conservan tal cual, como en el Excel original)."""
    if not tipo:
        return texto
    v = de_texto(texto, tipo)
    return texto if isinstance(v, str) or a_texto(v, tipo) != texto else v


def construir_libro(matrices: dict[str, Matriz]) -> openpyxl.Workbook:
    """Matrices de texto -> libro de openpyxl con los tipos de cada columna (fechas, enteros, casillas)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, m in matrices.items():
        ws = wb.create_sheet(nombre[:31])
        if not m:
            continue
        tipos = [_tipo_columna(nombre, h) for h in m[0]]
        for r, fila in enumerate(m, start=1):
            for c, texto in enumerate(fila, start=1):
                if texto == "":
                    continue
                ws.cell(row=r, column=c).value = texto if r == 1 else _valor_celda(texto, tipos[c - 1])
    return wb


def _registrar_variables(matrices: dict[str, Matriz]) -> None:
    """Deja registradas en el esquema las variables propias que define la hoja «Variables»."""
    m = matrices.get(S.HOJA_VARIABLES) or []
    filas = [dict(zip(m[0], f)) for f in m[1:]] if m else []
    S.registrar_variables(V.campos_de(V.desde_filas(filas)))


# ============================================================================ almacén
class SheetsStorage(AlmacenBase):
    """Libro de trabajo guardado en una hoja de cálculo de Google."""

    def __init__(self, cliente: ClienteHojas, semilla: Path | None = None, inicial: Path | None = None, ttl: float = 8.0):
        self.cliente = cliente
        self.semilla = Path(semilla) if semilla else None
        self.inicial = Path(inicial) if inicial else None      # libro local con el que llenar una hoja vacía
        self.ttl = ttl
        self._matrices: dict[str, Matriz] | None = None
        self._wb: openpyxl.Workbook | None = None
        self._marca = ""
        self._t_chequeo = 0.0
        self._lock = storage._LOCK

    # ---- identidad
    @property
    def url(self) -> str:
        return self.cliente.url

    @property
    def nombre(self) -> str:
        return "Google Sheets"

    @property
    def descripcion(self) -> str:
        return f"una hoja de cálculo de Google ([abrirla]({self.url}))" if self.url else "una hoja de cálculo de Google"

    # ---- sincronización con la hoja
    def _descargar(self, marca: str) -> None:
        crudas = self.cliente.leer()
        self._matrices = {k: normalizar(v) for k, v in crudas.items()}
        _registrar_variables(self._matrices)
        self._wb = construir_libro(self._matrices)
        self._marca = marca

    def _sincronizar(self, forzar: bool = False) -> None:
        t = time.monotonic()
        if self._matrices is not None and not forzar and t - self._t_chequeo < self.ttl:
            return
        marca = self.cliente.marca()
        self._t_chequeo = t
        if self._matrices is None or marca != self._marca:
            self._descargar(marca)

    def firma(self) -> str:
        with self._lock:
            self._sincronizar()
            return self._marca

    # ---- ciclo de vida
    def asegurar(self) -> None:
        """Si la hoja está vacía la llena con el libro local o el Excel original; si no, la migra al esquema vigente."""
        with self._lock:
            self._sincronizar(forzar=True)
            if S.HOJA_BASE not in (self._matrices or {}):
                origen = next((p for p in (self.inicial, self.semilla) if p and Path(p).exists()), None)
                if origen is None:
                    raise ErrorAlmacen("La hoja de Google está vacía y no hay un Excel original para llenarla.")
                wb = openpyxl.load_workbook(origen)
                _migrar(wb)
                self._publicar_completo(wb)
                return
            wb = construir_libro(self._matrices)
            if _migrar(wb):
                self._guardar(wb)

    # ---- lectura
    def leer_hoja_texto(self, nombre: str) -> pd.DataFrame:
        with self._lock:
            self._sincronizar()
            from utils.storage import df_de_hoja
            return df_de_hoja(self._wb, nombre)

    def leer_base_excel(self) -> pd.DataFrame:
        """Base_Datos con los tipos de Excel (fechas, enteros…), como la lee pandas en la vista de lectura."""
        with self._lock:
            self._sincronizar()
            solo = construir_libro({S.HOJA_BASE: self._matrices[S.HOJA_BASE]})
            buf = io.BytesIO()
            solo.save(buf)
        buf.seek(0)
        return pd.read_excel(buf, sheet_name=S.HOJA_BASE, engine="openpyxl")

    def exportar_xlsx(self) -> bytes:
        with self._lock:
            self._sincronizar()
            buf = io.BytesIO()
            construir_libro(self._matrices).save(buf)
            return buf.getvalue()

    # ---- escritura
    @contextmanager
    def transaccion(self):
        """Bloquea, lee la hoja vigente, entrega un `Libro` y sube solo lo que cambió al salir."""
        with self._lock:
            self._sincronizar(forzar=True)
            if S.HOJA_BASE not in self._matrices:
                raise ErrorAlmacen("La hoja de Google no tiene la pestaña de datos (Base_Datos).")
            wb = construir_libro(self._matrices)
            libro = Libro(wb)
            yield libro                       # si el cuerpo lanza una excepción, no se sube nada
            self._guardar(wb)

    def _guardar(self, wb: openpyxl.Workbook) -> None:
        nuevas = {ws.title: matriz_de_hoja(ws, ws.title) for ws in wb.worksheets}
        self._empujar(nuevas)
        self._matrices = nuevas
        self._wb = construir_libro(nuevas)
        self._marca = self.cliente.marca()
        self._t_chequeo = time.monotonic()

    def _empujar(self, nuevas: dict[str, Matriz]) -> None:
        """Sube a Google las pestañas nuevas y las filas que cambiaron en las existentes."""
        antes = self._matrices or {}
        for nombre, m in nuevas.items():
            if nombre not in antes:
                self.cliente.reemplazar(nombre, m)
                continue
            if m == antes[nombre]:
                continue
            viejo = antes[nombre]
            ancho = max((len(f) for f in m + viejo), default=0)
            pad = lambda f: list(f) + [""] * (ancho - len(f))  # noqa: E731
            if len(m) < len(viejo):                            # se quitaron filas: se reescribe la pestaña
                self.cliente.reemplazar(nombre, m)
                continue
            cambiadas = {i + 1: pad(f) for i, f in enumerate(m) if i >= len(viejo) or pad(f) != pad(viejo[i])}
            if len(cambiadas) > MAX_FILAS_PARCIAL:
                self.cliente.reemplazar(nombre, m)
            elif cambiadas:
                self.cliente.escribir_filas(nombre, cambiadas, ancho)

    def _publicar_completo(self, wb: openpyxl.Workbook) -> None:
        """Deja la hoja de Google idéntica al libro `wb` (primera carga o restablecer)."""
        nuevas = {ws.title: matriz_de_hoja(ws, ws.title) for ws in wb.worksheets}
        existentes = [t for t in self.cliente.titulos() if not t.startswith(PREFIJO_RESPALDO)]
        for nombre, m in nuevas.items():
            self.cliente.reemplazar(nombre, m)
        for t in existentes:                                  # pestañas que ya no existen (o la «Hoja 1» por defecto)
            if t not in nuevas and (t.lower() in HOJAS_POR_DEFECTO or t in (self._matrices or {})):
                self.cliente.borrar_hoja(t)
        self._matrices = nuevas
        _registrar_variables(nuevas)
        self._wb = construir_libro(nuevas)
        self._marca = self.cliente.marca()
        self._t_chequeo = time.monotonic()

    # ---- respaldos (pestañas RESPALDO_… con una copia de Base_Datos)
    def respaldar(self, etiqueta: str = "") -> str | None:
        with self._lock:
            self._sincronizar()
            base = (self._matrices or {}).get(S.HOJA_BASE)
            if not base:
                return None
            etq = re.sub(r"[^a-zA-Z0-9_-]+", "", etiqueta)
            nombre = f"{PREFIJO_RESPALDO}{datetime.now():%Y%m%d_%H%M%S}" + (f"_{etq}" if etq else "")
            self.cliente.reemplazar(nombre[:95], base)
            for viejo in self._nombres_respaldo()[storage.MAX_RESPALDOS:]:
                self.cliente.borrar_hoja(viejo)
            return nombre[:95]

    def _nombres_respaldo(self) -> list[str]:
        return sorted((t for t in self.cliente.titulos() if t.startswith(PREFIJO_RESPALDO)), reverse=True)

    def listar_respaldos(self) -> list[str]:
        return self._nombres_respaldo()

    def respaldos(self) -> list[Respaldo]:
        out = []
        for n in self._nombres_respaldo():
            m = re.match(rf"{PREFIJO_RESPALDO}(\d{{8}}_\d{{6}})", n)
            out.append(Respaldo(n, datetime.strptime(m.group(1), "%Y%m%d_%H%M%S") if m else None, None))
        return out

    def restablecer_desde_semilla(self) -> str | None:
        """Reemplaza el contenido de la hoja por el Excel original (guarda antes un respaldo)."""
        if self.semilla is None or not self.semilla.exists():
            raise ErrorAlmacen("No se encuentra el Excel original.")
        with self._lock:
            respaldo = self.respaldar("antes_de_restablecer")
            wb = openpyxl.load_workbook(self.semilla)
            _migrar(wb)
            self._publicar_completo(wb)
            return respaldo
