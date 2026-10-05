# -*- coding: utf-8 -*-
"""Almacenamiento del libro de trabajo.

La base de datos vive en una hoja de cálculo de Google (`utils/sheets.py`, clase `SheetsStorage`) cuando los
secretos de Streamlit la configuran. Sin esa configuración (desarrollo, pruebas) se usa un archivo local
`data/catalogo_gestion.xlsx` (`LocalStorage`), creado la primera vez copiando el Excel original, que queda
intacto como semilla. Ambos almacenes tienen la misma interfaz (`firma`, `transaccion`, `leer_hoja_texto`...)
y se les agregan las mismas columnas y hojas de gestión.

Reglas de diseño:
- Todas las escrituras ocurren dentro de `transaccion()`: bloqueo compartido, relectura del
  archivo vigente, y guardado atómico (archivo temporal + os.replace). Si algo falla dentro de
  la transacción no se guarda nada.
- Las celdas se comparan y registran en su forma de texto canónica (utils/validation.py), y al
  escribir se usa el tipo de la columna (fecha, entero...). Solo se tocan las celdas modificadas.
- Este módulo no importa Streamlit.
"""
from __future__ import annotations

import os
import re
import shutil
import threading
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import io
from dataclasses import dataclass

import openpyxl
import pandas as pd

from utils import schema as S
from utils import variables as V
from utils.validation import a_texto, de_texto, parse_fecha

RAIZ = Path(__file__).resolve().parent.parent
DIR_DATOS = RAIZ / "data"
RUTA_SEMILLA = DIR_DATOS / "experiencia_glocal_catálogo_web_histórico_hasta_10_Agosto_2026.xlsx"
RUTA_TRABAJO = Path(os.environ.get("IMPACTO_RUTA_TRABAJO", DIR_DATOS / "catalogo_gestion.xlsx"))
MAX_RESPALDOS = 5

# Un único bloqueo para todo el proceso: protege un archivo compartido por todas las sesiones.
_LOCK = threading.RLock()
_ALMACENES: dict[str, "LocalStorage"] = {}


def ahora() -> str:
    """Marca de tiempo local, en texto ISO sin microsegundos."""
    return datetime.now().isoformat(timespec="seconds")


class ErrorAlmacen(Exception):
    """Error de lectura/escritura del libro, con mensaje apto para mostrar al usuario."""


@dataclass
class Respaldo:
    """Copia de seguridad del libro (archivo o pestaña) para mostrar en Administración."""
    nombre: str
    fecha: datetime | None
    kb: int | None = None


def df_de_hoja(wb: openpyxl.Workbook, nombre: str) -> pd.DataFrame:
    """Hoja de un libro abierto en texto canónico (vacío = ''), sin filas totalmente vacías."""
    if nombre not in wb.sheetnames:
        return pd.DataFrame(columns=list(S.HOJAS_AUX.get(nombre, ())), dtype=object)
    filas = wb[nombre].iter_rows(values_only=True)
    cab = next(filas, None)
    if cab is None:
        return pd.DataFrame()
    nombres = [str(h) if h is not None else "" for h in cab]
    tipos = [(S.CAMPO[h].tipo if nombre == S.HOJA_BASE and h in S.CAMPO else None) for h in nombres]
    datos = []
    for fila in filas:
        textos = [a_texto(v, t) for v, t in zip(fila, tipos)]
        if any(textos):
            datos.append(textos + [""] * (len(nombres) - len(textos)))
    df = pd.DataFrame(datos, columns=nombres, dtype=object)
    return df.loc[:, [c for c in df.columns if c]]


class AlmacenBase:
    """Lecturas comunes a todos los almacenes (se apoyan en `leer_hoja_texto`)."""

    def leer_variables(self) -> list[V.Variable]:
        """Variables propias del libro; además las registra en el esquema."""
        df = self.leer_hoja_texto(S.HOJA_VARIABLES)
        vars_ = V.desde_filas(df.to_dict("records"))
        S.registrar_variables(V.campos_de(vars_))
        return vars_

    def leer_base_texto(self) -> pd.DataFrame:
        """Base_Datos cruda en forma de texto canónico, con todas las columnas del esquema (y las variables propias)."""
        self.leer_variables()                      # antes: así las columnas var_* se leen con su tipo
        df = self.leer_hoja_texto(S.HOJA_BASE)
        for c in S.COLUMNAS:
            if c not in df.columns:
                df[c] = ""
        return df.reset_index(drop=True)


# ============================================================================ handle de libro
class Libro:
    """Vista de trabajo sobre un libro abierto con openpyxl (solo dentro de una transacción)."""

    def __init__(self, wb: openpyxl.Workbook):
        self.wb = wb
        self.ws = wb[S.HOJA_BASE]
        self._cols: dict[str, int] = {}
        self._filas: dict[str, int] = {}
        self._reindexar()
        self.sincronizar_variables()

    # ---- índices
    def _reindexar(self) -> None:
        self._cols = {}
        for i, c in enumerate(self.ws[1], start=1):
            if c.value not in (None, ""):
                self._cols[str(c.value)] = i
        self._filas = {}
        col_id = self._cols.get("id_evento")
        if col_id:
            for r in range(2, self.ws.max_row + 1):
                v = self.ws.cell(row=r, column=col_id).value
                if v not in (None, ""):
                    self._filas[str(v).strip()] = r

    def ids(self) -> list[str]:
        return list(self._filas)

    def existe(self, id_evento: str) -> bool:
        return id_evento in self._filas

    @property
    def columnas(self) -> list[str]:
        return list(self._cols)

    def _tipo(self, campo: str) -> str | None:
        c = S.CAMPO.get(campo)
        return c.tipo if c else None

    # ---- lectura (texto canónico)
    def valor(self, id_evento: str, campo: str) -> str:
        r, c = self._filas.get(id_evento), self._cols.get(campo)
        if r is None or c is None:
            return ""
        return a_texto(self.ws.cell(row=r, column=c).value, self._tipo(campo))

    def fila(self, id_evento: str) -> dict[str, str]:
        return {k: self.valor(id_evento, k) for k in self._cols}

    # ---- escritura
    def set_valor(self, id_evento: str, campo: str, texto: str) -> None:
        r = self._filas[id_evento]
        c = self._cols.get(campo)
        if c is None:
            c = self._agregar_columna(campo)
        self.ws.cell(row=r, column=c).value = de_texto(texto, self._tipo(campo) or S.TEXTO)

    def _agregar_columna(self, campo: str) -> int:
        c = len(self._cols) + 1
        self.ws.cell(row=1, column=c).value = campo
        self._cols[campo] = c
        return c

    def asegurar_columna(self, campo: str) -> int:
        """Índice de la columna `campo` en Base_Datos; la crea al final si no existe."""
        return self._cols.get(campo) or self._agregar_columna(campo)

    def variables(self) -> list[V.Variable]:
        """Variables propias definidas en este libro (hoja Variables)."""
        if S.HOJA_VARIABLES not in self.wb.sheetnames:
            return []
        return V.desde_filas([f for _, f in self.registros(S.HOJA_VARIABLES)])

    def sincronizar_variables(self) -> list[V.Variable]:
        """Registra en el esquema las variables propias de este libro (para validar, tipar y mostrar sus columnas)."""
        vars_ = self.variables()
        S.registrar_variables(V.campos_de(vars_))
        return vars_

    def siguiente_id(self) -> str:
        n = 0
        for ide in self._filas:
            m = re.search(r"(\d+)$", ide)
            if m:
                n = max(n, int(m.group(1)))
        return f"EV{n + 1:04d}"

    def agregar_fila(self, datos: dict[str, str]) -> str:
        """Agrega un registro al final de Base_Datos. `datos` debe traer id_evento."""
        ide = datos["id_evento"]
        r = self.ws.max_row + 1
        for campo, texto in datos.items():
            c = self._cols.get(campo) or self._agregar_columna(campo)
            self.ws.cell(row=r, column=c).value = de_texto(texto, self._tipo(campo) or S.TEXTO)
        self._filas[ide] = r
        return ide

    def eliminar_fila(self, id_evento: str) -> None:
        self.ws.delete_rows(self._filas[id_evento])
        self._reindexar()

    # ---- hojas auxiliares (todo texto)
    def hoja(self, nombre: str):
        if nombre not in self.wb.sheetnames:
            ws = self.wb.create_sheet(nombre)
            for i, h in enumerate(S.HOJAS_AUX[nombre], start=1):
                ws.cell(row=1, column=i).value = h
        return self.wb[nombre]

    def registros(self, nombre: str) -> list[tuple[int, dict[str, str]]]:
        """[(n.º de fila, {columna: texto})] de una hoja auxiliar."""
        ws = self.hoja(nombre)
        cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
        out = []
        for r in range(2, ws.max_row + 1):
            fila = {cab[i]: a_texto(ws.cell(row=r, column=i + 1).value) for i in range(len(cab)) if cab[i]}
            if any(fila.values()):
                out.append((r, fila))
        return out

    def agregar_registro(self, nombre: str, datos: dict[str, str]) -> None:
        ws = self.hoja(nombre)
        cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
        r = ws.max_row + 1
        for i, h in enumerate(cab, start=1):
            if h and h in datos:
                ws.cell(row=r, column=i).value = a_texto(datos[h]) or None

    def set_registro(self, nombre: str, n_fila: int, campo: str, texto: str) -> None:
        ws = self.hoja(nombre)
        cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
        ws.cell(row=n_fila, column=cab.index(campo) + 1).value = a_texto(texto) or None

    def max_numero(self, nombre: str, columna: str) -> int:
        """Mayor número al final de los valores de una columna (p. ej. 'CH000123' -> 123); 0 si no hay."""
        ws = self.hoja(nombre)
        cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
        if columna not in cab:
            return 0
        col = cab.index(columna) + 1
        n = 0
        for r in range(2, ws.max_row + 1):
            m = re.search(r"(\d+)$", a_texto(ws.cell(row=r, column=col).value))
            if m:
                n = max(n, int(m.group(1)))
        return n


# ============================================================================ migración
def _migrar(wb: openpyxl.Workbook) -> bool:
    """Deja el libro con las columnas y hojas de gestión. Idempotente. True si cambió algo."""
    cambio = False
    ws = wb[S.HOJA_BASE]
    cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
    nuevas = [c for c in S.COLUMNAS_NUEVAS if c not in cab]
    if nuevas:
        col0 = len([h for h in cab if h]) + 1
        for i, nombre in enumerate(nuevas):
            ws.cell(row=1, column=col0 + i).value = nombre
        cab = [str(c.value) if c.value is not None else "" for c in ws[1]]
        idx = {h: i + 1 for i, h in enumerate(cab) if h}
        c_ide, c_ts = idx.get("id_evento"), idx.get("timestamp_extraccion")
        for r in range(2, ws.max_row + 1):
            if c_ide and ws.cell(row=r, column=c_ide).value in (None, ""):
                continue
            if "origen" in nuevas:
                ws.cell(row=r, column=idx["origen"]).value = "historico"
            if "fecha_carga" in nuevas and c_ts:
                ts = parse_fecha(ws.cell(row=r, column=c_ts).value)
                ws.cell(row=r, column=idx["fecha_carga"]).value = ts
        cambio = True

    libro = Libro(wb)
    for nombre in S.HOJAS_AUX:
        if nombre not in wb.sheetnames:
            libro.hoja(nombre)
            cambio = True
            if nombre == S.HOJA_CATEGORIAS:
                _sembrar_categorias(libro)

    if nuevas or S.HOJA_LIBRO not in wb.sheetnames:
        from utils import libro_codigos  # import perezoso: evita ciclo con este módulo
        libro_codigos.regenerar(wb)
        cambio = True
    return cambio


def _sembrar_categorias(libro: Libro) -> None:
    """Catálogo inicial de categorías a partir de los valores que ya usa la base."""
    ts = ahora()
    for dim in S.DIMENSIONES_CATALOGO:
        vistos: dict[str, None] = {}
        for ide in libro.ids():
            for p in S.dividir_etiquetas(libro.valor(ide, dim)):
                if p.lower() not in S.VALORES_NO_ETIQUETA:
                    vistos[p] = None
        for nombre in sorted(vistos, key=str.casefold):
            libro.agregar_registro(S.HOJA_CATEGORIAS, {
                "dimension": dim, "nombre": nombre, "descripcion": "", "activa": "True",
                "creada_por": "(migración)", "fecha_hora": ts,
            })


# ============================================================================ almacén local
class LocalStorage(AlmacenBase):
    url = ""                                   # solo el almacén de Google Sheets tiene dirección web

    def __init__(self, ruta: Path, semilla: Path | None = None):
        self.ruta = Path(ruta)
        self.semilla = Path(semilla) if semilla else None

    @property
    def nombre(self) -> str:
        return self.ruta.name

    @property
    def descripcion(self) -> str:
        return f"el archivo local `{self.ruta.name}` (solo para pruebas y desarrollo)"

    def exportar_xlsx(self) -> bytes:
        """Copia completa del libro como archivo .xlsx."""
        return self.ruta.read_bytes()

    def leer_base_excel(self) -> pd.DataFrame:
        """Base_Datos tal como la lee pandas del libro (tipos de Excel), para la vista de lectura."""
        return pd.read_excel(self.ruta, sheet_name=S.HOJA_BASE, engine="openpyxl")

    # ---- ciclo de vida
    def asegurar(self) -> None:
        """Crea el archivo de trabajo si no existe y lo migra al esquema vigente."""
        with _LOCK:
            if not self.ruta.exists():
                if self.semilla is None or not self.semilla.exists():
                    raise ErrorAlmacen(
                        "No existe el archivo de trabajo ni el Excel original para crearlo "
                        f"({self.ruta.name})."
                    )
                self.ruta.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(self.semilla, self.ruta)
            wb = openpyxl.load_workbook(self.ruta)
            if _migrar(wb):
                self._guardar(wb)

    def firma(self) -> str:
        """Cambia cada vez que el archivo se modifica; sirve de clave de caché."""
        try:
            st = self.ruta.stat()
        except FileNotFoundError:
            return "ausente"
        return f"{st.st_mtime_ns}-{st.st_size}"

    def _guardar(self, wb: openpyxl.Workbook) -> None:
        tmp = self.ruta.with_name(self.ruta.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, self.ruta)

    # ---- lectura
    def _abrir_lectura(self):
        try:
            return openpyxl.load_workbook(self.ruta, read_only=True, data_only=True)
        except FileNotFoundError as e:
            raise ErrorAlmacen(f"No se encuentra el archivo de datos ({self.ruta.name}).") from e
        except PermissionError as e:
            raise ErrorAlmacen(
                "El archivo de datos está abierto en otro programa. Ciérralo e inténtalo de nuevo."
            ) from e

    def leer_hoja_texto(self, nombre: str) -> pd.DataFrame:
        """Hoja completa en forma de texto canónico (vacío = ''). Sin filas totalmente vacías."""
        with _LOCK:
            wb = self._abrir_lectura()
            try:
                return df_de_hoja(wb, nombre)
            finally:
                wb.close()

    # ---- escritura
    @contextmanager
    def transaccion(self):
        """Bloquea, abre el archivo vigente, entrega un `Libro` y guarda atómicamente al salir."""
        with _LOCK:
            try:
                wb = openpyxl.load_workbook(self.ruta)
            except FileNotFoundError as e:
                raise ErrorAlmacen(f"No se encuentra el archivo de datos ({self.ruta.name}).") from e
            except PermissionError as e:
                raise ErrorAlmacen(
                    "El archivo de datos está abierto en otro programa. Ciérralo e inténtalo de nuevo."
                ) from e
            libro = Libro(wb)
            yield libro          # si el cuerpo lanza una excepción, no se guarda nada
            try:
                self._guardar(wb)
            except PermissionError as e:
                raise ErrorAlmacen(
                    "No se pudo guardar: el archivo de datos está abierto en otro programa."
                ) from e

    # ---- respaldos
    def _patron_respaldo(self) -> str:
        return f"{self.ruta.stem}.BACKUP_*{self.ruta.suffix}"

    def respaldar(self, etiqueta: str = "") -> Path | None:
        """Copia el archivo con marca de tiempo y conserva solo los últimos MAX_RESPALDOS."""
        with _LOCK:
            if not self.ruta.exists():
                return None
            etq = re.sub(r"[^a-zA-Z0-9_-]+", "", etiqueta)
            nombre = f"{self.ruta.stem}.BACKUP_{datetime.now():%Y%m%d_%H%M%S}"
            destino = self.ruta.with_name(f"{nombre}{'_' + etq if etq else ''}{self.ruta.suffix}")
            shutil.copy2(self.ruta, destino)
            for viejo in self.listar_respaldos()[MAX_RESPALDOS:]:
                try:
                    viejo.unlink()
                except OSError:
                    pass
            return destino

    def listar_respaldos(self) -> list[Path]:
        return sorted(self.ruta.parent.glob(self._patron_respaldo()), reverse=True)

    def respaldos(self) -> list[Respaldo]:
        return [Respaldo(p.name, datetime.fromtimestamp(p.stat().st_mtime), round(p.stat().st_size / 1024))
                for p in self.listar_respaldos()]

    def restablecer_desde_semilla(self) -> Path | None:
        """Reemplaza el libro por una copia nueva del Excel original (guarda antes un respaldo)."""
        if self.semilla is None or not self.semilla.exists():
            raise ErrorAlmacen("No se encuentra el Excel original.")
        with _LOCK:
            respaldo = self.respaldar("antes_de_restablecer")
            shutil.copy2(self.semilla, self.ruta)
            wb = openpyxl.load_workbook(self.ruta)
            _migrar(wb)
            self._guardar(wb)
            return respaldo


# ============================================================================ acceso
def get_storage():
    """Almacén de trabajo de este proceso, ya creado y migrado: Google Sheets si los secretos lo configuran,
    y si no, el archivo local."""
    with _LOCK:
        alm = _ALMACENES.get("trabajo")
        if alm is None:
            from utils import sheets  # import perezoso: sheets.py importa este módulo
            cfg = sheets.configuracion()
            alm = sheets.crear(cfg, RUTA_SEMILLA, RUTA_TRABAJO) if cfg else LocalStorage(RUTA_TRABAJO, RUTA_SEMILLA)
            alm.asegurar()
            _ALMACENES["trabajo"] = alm
        return alm


def olvidar_almacenes() -> None:
    """Para pruebas: descarta los almacenes en memoria."""
    with _LOCK:
        _ALMACENES.clear()
