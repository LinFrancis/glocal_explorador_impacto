# -*- coding: utf-8 -*-
"""Generador de la hoja `Libro_de_Codigos` a partir del registro de campos.

Conserva lo que ya está escrito a mano (descripción, tipo, opciones y fuente de cada columna),
agrega las columnas nuevas de la plataforma y reemplaza las cifras fijas de las opciones de
respuesta ("~577 instituciones", "139 cuencas"...) por recuentos calculados desde los datos.
"""
from __future__ import annotations

import openpyxl

from utils import schema as S
from utils.validation import a_texto

COLUMNAS_LIBRO = ("columna", "descripcion", "tipo_variable", "opciones_respuesta", "fuente")

_FUENTE_PLATAFORMA = "Plataforma de gestión Impacto Glocal"

# Textos de las columnas que agrega la plataforma (el resto ya está en el libro original).
_NUEVAS = {
    "wp_id": (
        "Identificador numérico de la entrada en WordPress. Permite reconocer una noticia aunque "
        "cambie su URL o su slug.",
        "Numérica (entero)", "Vacío si el registro no vino de la web",
        "API REST de WordPress; completado por la plataforma de gestión",
    ),
    "origen": (
        "Cómo llegó el registro a la base: histórico (catálogo original), scraping (sincronización "
        "con la web), manual (formulario) o archivo (importación de Excel/CSV).",
        "Categórica única", "historico; scraping; manual; archivo", _FUENTE_PLATAFORMA,
    ),
    "cargado_por": (
        "Nombre de la persona que cargó el registro (el que eligió de la lista del equipo al iniciar sesión).",
        "Texto libre", "Vacío en los registros históricos", _FUENTE_PLATAFORMA,
    ),
    "fecha_carga": (
        "Fecha y hora en que el registro se cargó a la base. En los registros históricos es la "
        "fecha de extracción original.",
        "Fecha-hora (ISO 8601)", "", _FUENTE_PLATAFORMA,
    ),
    "editado_por": (
        "Nombre de la última persona que modificó el registro.",
        "Texto libre", "Vacío si nunca se ha editado", _FUENTE_PLATAFORMA,
    ),
    "fecha_edicion": (
        "Fecha y hora de la última modificación del registro.",
        "Fecha-hora (ISO 8601)", "Vacío si nunca se ha editado", _FUENTE_PLATAFORMA,
    ),
    "carpeta_proyecto": (
        "Enlace a la carpeta del proyecto (Google Drive u otro lugar con sus archivos).",
        "URL", "Vacío si aún no se indica", _FUENTE_PLATAFORMA,
    ),
    "documentos_proyecto": (
        "Otros enlaces del proyecto: archivos finales, documentación, informes.",
        "Lista de URLs (separador ' | ')", "Vacío si aún no se indican", _FUENTE_PLATAFORMA,
    ),
}

_FUENTE_VARIABLE = "Variable propia definida en Administración (plataforma de gestión)"


def _tipo_de_variable(c) -> str:
    if c.tipo == S.OPCION:
        return "Categórica única"
    if c.tipo == S.ETIQUETAS:
        return "Categórica multi-etiqueta (';')"
    return S.TIPO_NOMBRE.get(c.tipo, c.tipo)

# Opciones de respuesta con recuento calculado: columna -> plantilla ({n} = valores distintos).
_PLANTILLAS = {
    "categorias": "{n} categorías temáticas (ver Glosario en la app)",
    "metodologia": "{n} metodologías + 'No especificado' (ver Glosario)",
    "actores_normalizados": "Abierto ({n} instituciones canónicas)",
    "objetivo_gcaa": "{n} objetivos numerados + 'No aplica'",
    "sitios_cuenca_nombre": "{n} cuencas de Chile o vacío",
}


def _distintos(filas: list[dict], col: str, excluir: set[str]) -> int:
    """Valores distintos de una columna multi-etiqueta (unificando los alias de metodología,
    igual que la vista de lectura de la app)."""
    vistos: set[str] = set()
    for f in filas:
        for p in S.dividir_etiquetas(f.get(col, "")):
            if p.lower() not in excluir:
                vistos.add(S.ALIAS_METODOLOGIA.get(p.lower(), p) if col == "metodologia" else p)
    return len(vistos)


def _n_cuencas_referencia() -> int | None:
    """Cuencas únicas (por código) de la tabla de referencia BNA/DGA, si está disponible."""
    try:
        import pandas as pd
        from utils.storage import DIR_DATOS
        ref = pd.read_excel(DIR_DATOS / "cuencas_chile_bna.xlsx", sheet_name="Cuencas")
        return int(ref["COD_CUEN"].nunique())
    except Exception:  # noqa: BLE001  (el recuento es accesorio: si falla se conserva el texto)
        return None


def regenerar(wb: openpyxl.Workbook) -> None:
    """Reescribe la hoja Libro_de_Codigos del libro abierto `wb`."""
    ws_base = wb[S.HOJA_BASE]
    cab = [str(c.value) if c.value is not None else "" for c in ws_base[1]]
    filas = []
    for r in range(2, ws_base.max_row + 1):
        valores = [ws_base.cell(row=r, column=i + 1).value for i in range(len(cab))]
        if any(v not in (None, "") for v in valores):
            filas.append({cab[i]: a_texto(valores[i]) for i in range(len(cab)) if cab[i]})

    previo: dict[str, dict[str, str]] = {}
    if S.HOJA_LIBRO in wb.sheetnames:
        ws_old = wb[S.HOJA_LIBRO]
        head = [str(c.value) if c.value is not None else "" for c in ws_old[1]]
        for r in range(2, ws_old.max_row + 1):
            reg = {head[i]: a_texto(ws_old.cell(row=r, column=i + 1).value) for i in range(len(head)) if head[i]}
            if reg.get("columna"):
                previo[reg["columna"]] = reg

    excluir = S.VALORES_NO_ETIQUETA
    n_cuencas = _n_cuencas_referencia()
    nuevas_filas: list[dict[str, str]] = []
    for c in S.CAMPOS:
        reg = dict(previo.get(c.key, {}))
        if not reg:
            if c.key in _NUEVAS:
                d, t, o, f = _NUEVAS[c.key]
                reg = {"descripcion": d, "tipo_variable": t, "opciones_respuesta": o, "fuente": f}
            else:
                reg = {"descripcion": c.label, "tipo_variable": S.TIPO_NOMBRE.get(c.tipo, ""),
                       "opciones_respuesta": "", "fuente": _FUENTE_PLATAFORMA}
        if c.variable:                                   # las variables propias siempre reflejan su definición actual
            opciones = "; ".join(c.opciones) if isinstance(c.opciones, tuple) else ""
            reg = {"descripcion": c.ayuda or f"Variable propia «{c.label}».", "tipo_variable": _tipo_de_variable(c),
                   "opciones_respuesta": opciones, "fuente": _FUENTE_VARIABLE}
        if c.key in _PLANTILLAS:
            n = n_cuencas if c.key == "sitios_cuenca_nombre" else _distintos(filas, c.key, excluir)
            if n:
                reg["opciones_respuesta"] = _PLANTILLAS[c.key].format(n=n)
        reg["columna"] = c.key
        nuevas_filas.append(reg)
    # Columnas que estaban en el libro pero ya no están en el registro: se conservan al final.
    registradas = {c.key for c in S.CAMPOS}
    nuevas_filas += [reg for k, reg in previo.items() if k not in registradas]

    if S.HOJA_LIBRO in wb.sheetnames:
        ws = wb[S.HOJA_LIBRO]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(S.HOJA_LIBRO)
    for i, h in enumerate(COLUMNAS_LIBRO, start=1):
        ws.cell(row=1, column=i).value = h
    for r, reg in enumerate(nuevas_filas, start=2):
        for i, h in enumerate(COLUMNAS_LIBRO, start=1):
            ws.cell(row=r, column=i).value = reg.get(h, "") or None
