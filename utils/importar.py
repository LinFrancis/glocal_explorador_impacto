# -*- coding: utf-8 -*-
"""Importación de noticias desde un archivo Excel/CSV: lectura, correspondencia de columnas,
evaluación fila a fila (errores, avisos, duplicados) y plantilla descargable.

Lógica pura (sin Streamlit): la página de carga solo la dibuja.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from urllib.parse import urlparse

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from utils import dedupe
from utils import schema as S
from utils.validation import a_texto, validar_fila

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 2000

# Columnas de la plantilla: lo elemental + lo que más se completa al cargar.
PLANTILLA_COLUMNAS = (
    "titulo", "url_noticia", "fuente", "fecha_publicacion_web", "descripcion_catalogo",
    "contenido_completo", "imagen_principal_url", "tipo_informacion", "categoria_macro", "categorias",
    "metodologia", "actores_normalizados", "lugar", "eje_gcaa", "atributos_resiliencia",
    "beneficiarios_directos", "enfoque_genero", "carpeta_proyecto",
)

# Nombres alternativos habituales de las columnas (ya normalizados con _norm).
_ALIAS = {
    "url": "url_noticia", "enlace": "url_noticia", "link": "url_noticia", "direccion": "url_noticia",
    "resumen": "descripcion_catalogo", "descripcion": "descripcion_catalogo",
    "texto": "contenido_completo", "contenido": "contenido_completo",
    "fecha": "fecha_publicacion_web", "fechadepublicacion": "fecha_publicacion_web",
    "imagen": "imagen_principal_url", "foto": "imagen_principal_url",
    "categoria": "categorias", "categoriatematica": "categorias",
    "actores": "actores_normalizados", "instituciones": "actores_normalizados",
    "sitio": "fuente", "medio": "fuente",
}

ESTADO_NUEVA = "Nueva"
ESTADO_AVISOS = "Con avisos"
ESTADO_ERROR = "Con errores"
ESTADO_DUPLICADA = "Ya existe"
ESTADO_SOSPECHOSA = "Posible duplicada"


class ErrorImportacion(Exception):
    """Problema al leer el archivo; el mensaje se muestra tal cual al usuario."""


def _norm(texto) -> str:
    s = unicodedata.normalize("NFKD", str(texto or "")).lower()
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", s)


def _indice_de_nombres() -> dict[str, str]:
    indice: dict[str, str] = {}
    for c in S.CAMPOS:
        indice[_norm(c.key)] = c.key
        indice[_norm(c.label)] = c.key
    for alias, key in _ALIAS.items():
        indice.setdefault(alias, key)
    return indice


def mapear_columnas(encabezados: list[str]) -> dict[str, str]:
    """{encabezado del archivo: campo de la base} para los que se reconocen (sin repetir campos)."""
    indice = _indice_de_nombres()
    resultado: dict[str, str] = {}
    usados: set[str] = set()
    for h in encabezados:
        key = indice.get(_norm(h))
        if key and key not in usados and key not in ("id_evento",) + S.COLUMNAS_SISTEMA:
            resultado[h] = key
            usados.add(key)
    return resultado


def _separador_csv(texto: str) -> str:
    """Coma, punto y coma, tabulación o barra vertical (el detector genérico falla con una sola columna)."""
    muestra = "\n".join(texto.splitlines()[:20])
    try:
        return csv.Sniffer().sniff(muestra, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def hojas_de(datos: bytes, nombre: str) -> list[str]:
    if not nombre.lower().endswith((".xlsx", ".xlsm")):
        return []
    try:
        return pd.ExcelFile(io.BytesIO(datos)).sheet_names
    except Exception as e:  # noqa: BLE001
        raise ErrorImportacion("No se pudo abrir el Excel. ¿Está dañado o protegido con contraseña?") from e


def leer_archivo(datos: bytes, nombre: str, hoja: str | None = None) -> pd.DataFrame:
    """Lee un .xlsx o .csv como texto (todo `str`, vacío = ''). Lanza ErrorImportacion si no se puede."""
    if len(datos) > MAX_BYTES:
        raise ErrorImportacion(f"El archivo pesa más de {MAX_BYTES // (1024 * 1024)} MB. Divídelo en partes.")
    bajo = nombre.lower()
    try:
        if bajo.endswith((".xlsx", ".xlsm")):
            df = pd.read_excel(io.BytesIO(datos), sheet_name=hoja or 0, dtype=str)
        elif bajo.endswith(".csv"):
            texto = None
            for cod in ("utf-8-sig", "latin-1"):
                try:
                    texto = datos.decode(cod)
                    break
                except UnicodeDecodeError:
                    continue
            df = pd.read_csv(io.StringIO(texto), sep=_separador_csv(texto), dtype=str)
        else:
            raise ErrorImportacion("Formato no admitido: usa un archivo .xlsx o .csv.")
    except ErrorImportacion:
        raise
    except Exception as e:  # noqa: BLE001
        raise ErrorImportacion(f"No se pudo leer el archivo ({type(e).__name__}). Revisa que no esté dañado.") from e
    df = df.fillna("").astype(str)
    df = df.loc[:, [c for c in df.columns if str(c).strip() and not str(c).startswith("Unnamed")]]
    df = df[(df != "").any(axis=1)].reset_index(drop=True)
    if df.empty or df.shape[1] == 0:
        raise ErrorImportacion("El archivo no tiene filas con datos.")
    if len(df) > MAX_FILAS:
        raise ErrorImportacion(f"El archivo tiene {len(df)} filas; el máximo por importación es {MAX_FILAS}.")
    return df


def inferir_fuente(url: str) -> str:
    host = urlparse(url if "://" in url else "https://" + url).netloc.lower().removeprefix("www.")
    return next((f for f in S.FUENTES if host == f or host.endswith("." + f)), "")


def filas_desde_df(df: pd.DataFrame, mapa: dict[str, str]) -> list[dict[str, str]]:
    """Registros (campo -> texto) a partir del DataFrame y la correspondencia de columnas."""
    filas = []
    for r in df.to_dict("records"):
        fila = {campo: a_texto(r.get(h, "")) for h, campo in mapa.items()}
        if not fila.get("fuente") and fila.get("url_noticia"):
            fila["fuente"] = inferir_fuente(fila["url_noticia"])
        filas.append(fila)
    return filas


@dataclass
class Evaluacion:
    n: int                                   # n.º de fila en el archivo (desde 1)
    fila: dict[str, str]                     # ya normalizada
    estado: str
    errores: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    coincidencias: list[dedupe.Coincidencia] = field(default_factory=list)

    @property
    def importable(self) -> bool:
        return self.estado in (ESTADO_NUEVA, ESTADO_AVISOS)

    @property
    def detalle(self) -> str:
        partes = list(self.errores)
        partes += [f"{c.motivo} ({c.id_evento})" for c in self.coincidencias[:2]]
        partes += self.avisos[:2]
        return " · ".join(partes)


def evaluar(filas: list[dict[str, str]], indice_base: dedupe.Indice) -> list[Evaluacion]:
    """Clasifica cada fila: nueva, con avisos, con errores, ya existe o posible duplicada.

    Los repetidos DENTRO del archivo cuentan como duplicados de la primera aparición.
    """
    vistos = dedupe.Indice()
    salida: list[Evaluacion] = []
    for i, bruta in enumerate(filas, start=1):
        fila, errores, avisos = validar_fila(bruta)
        ev = Evaluacion(n=i, fila=fila, estado=ESTADO_NUEVA, errores=errores, avisos=avisos)
        if errores:
            ev.estado = ESTADO_ERROR
        else:
            coincidencias = indice_base.buscar(fila) + [
                dedupe.Coincidencia(c.nivel, f"fila {c.id_evento}" if c.id_evento else "otra fila del archivo",
                                    c.motivo + " (en este archivo)", c.entre_fuentes)
                for c in vistos.buscar(fila)
            ]
            ev.coincidencias = coincidencias
            if any(c.segura for c in coincidencias):
                ev.estado = ESTADO_DUPLICADA
            elif coincidencias:
                ev.estado = ESTADO_SOSPECHOSA
            elif avisos:
                ev.estado = ESTADO_AVISOS
            vistos.agregar({**fila, "id_evento": str(i)})
        salida.append(ev)
    return salida


def _columnas_plantilla() -> list[str]:
    return list(PLANTILLA_COLUMNAS) + [c.key for c in S.variables_activas()]


def plantilla_xlsx() -> bytes:
    """Excel de ejemplo: hoja «Noticias» con los encabezados y una fila de muestra, y hoja «Instrucciones»."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Noticias"
    ejemplo = {
        "titulo": "Taller de gobernanza del agua en Valparaíso",
        "url_noticia": "https://glocalminds.com/catalogo/taller-gobernanza-agua/",
        "fuente": "glocalminds.com",
        "fecha_publicacion_web": "2026-09-15",
        "descripcion_catalogo": "Resumen breve de la experiencia.",
        "categoria_macro": "Territorio, medioambiente y sostenibilidad",
        "lugar": "Valparaíso, Chile",
    }
    columnas = _columnas_plantilla()
    for j, key in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=S.etiqueta(key) + (" *" if S.CAMPO[key].obligatorio else ""))
        c.font = Font(bold=True)
        ws.cell(row=2, column=j, value=ejemplo.get(key, None))
        ws.column_dimensions[c.column_letter].width = 28
    ins = wb.create_sheet("Instrucciones")
    for j, h in enumerate(("Columna", "Obligatoria", "Tipo", "Ayuda"), start=1):
        ins.cell(row=1, column=j, value=h).font = Font(bold=True)
    for i, key in enumerate(columnas, start=2):
        c = S.CAMPO[key]
        ins.cell(row=i, column=1, value=c.label)
        ins.cell(row=i, column=2, value="Sí" if c.obligatorio else "No")
        ins.cell(row=i, column=3, value=S.TIPO_NOMBRE.get(c.tipo, c.tipo))
        ayuda = c.ayuda or ""
        if c.variable and isinstance(c.opciones, tuple):
            ayuda = (ayuda + " " if ayuda else "") + "Opciones: " + "; ".join(c.opciones) + (" (se pueden elegir varias, separadas por «;»)" if c.tipo == S.ETIQUETAS else "")
        ins.cell(row=i, column=4, value=ayuda)
    ins.cell(row=len(columnas) + 3, column=1,
             value="Las columnas con * son obligatorias. Borra la fila de ejemplo antes de subir el archivo. "
                   "Varias etiquetas en una misma celda se separan con «;».")
    for col, ancho in zip("ABCD", (34, 12, 28, 70)):
        ins.column_dimensions[col].width = ancho
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
