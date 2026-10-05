# -*- coding: utf-8 -*-
"""Variables analíticas propias: definición, registro dinámico, valores, opciones y migración."""
import openpyxl
import pytest

from utils import repo
from utils import schema as S
from utils import variables as V
from utils.repo import Cambio, ErrorOperacion
from utils.storage import LocalStorage

TAMANO = ("Pequeño", "Mediano", "Grande")


def _crear(almacen, etiqueta="Tamaño del proyecto", tipo="opcion", opciones=TAMANO, desc="", usuario="Francis"):
    return repo.crear_variable(etiqueta, tipo, list(opciones), desc, usuario, almacen)


def _valor(almacen, ide, clave):
    base = almacen.leer_base_texto().set_index("id_evento", drop=False)
    return str(base.at[ide, clave])


# ============================================================================ lógica pura
def test_nombres_y_claves():
    assert V.slug("Tamaño del proyecto") == "tamano_del_proyecto"
    assert V.clave_de("Tamaño del proyecto") == "var_tamano_del_proyecto"
    assert V.slug("  ¿Cuántas personas?  ") == "cuantas_personas"
    assert V.limpiar_etiqueta("tamaño_proyecto") == "Tamaño proyecto"            # escrito como identificador
    assert V.limpiar_etiqueta("  Tamaño   del  proyecto ") == "Tamaño del proyecto"
    assert V.limpiar_etiqueta("mi_variable con espacios") == "mi_variable con espacios"


def test_opciones_se_limpian_y_validan():
    assert V.normalizar_opciones("  pequeño \n\nmediano\n grande ", "opcion") == ("pequeño", "mediano", "grande")
    assert V.normalizar_opciones(["A", "B"], "etiquetas") == ("A", "B")
    assert V.normalizar_opciones(["x", "y"], "texto") == () and V.normalizar_opciones([], "si_no") == ()
    with pytest.raises(V.ErrorVariable, match="al menos 2"):
        V.normalizar_opciones(["solo una"], "opcion")
    with pytest.raises(V.ErrorVariable, match="repetida"):
        V.normalizar_opciones(["Grande", "grande", "otra"], "opcion")
    with pytest.raises(V.ErrorVariable, match="repetida"):
        V.normalizar_opciones(["Pequeño", "pequeno", "otra"], "opcion")
    with pytest.raises(V.ErrorVariable, match="«;» ni «\\|»"):
        V.normalizar_opciones(["a;b", "c"], "opcion")
    with pytest.raises(V.ErrorVariable, match="demasiado larga"):
        V.normalizar_opciones(["x" * 80, "y"], "opcion")


@pytest.mark.parametrize("tipo,tipo_campo,opciones", [
    ("opcion", S.OPCION, TAMANO), ("etiquetas", S.ETIQUETAS, TAMANO), ("si_no", S.OPCION, ("Sí", "No")),
    ("texto", S.TEXTO, None), ("texto_largo", S.TEXTO_LARGO, None), ("numero", S.NUMERO, None),
    ("fecha", S.FECHA, None), ("url", S.URL, None),
])
def test_cada_tipo_se_convierte_en_un_campo(tipo, tipo_campo, opciones):
    v = V.Variable("var_x", "X", tipo, TAMANO if tipo in ("opcion", "etiquetas") else (), "ayuda", True, 0)
    c = V.a_campo(v)
    assert (c.key, c.label, c.tipo, c.grupo, c.variable, c.activo, c.ayuda) == ("var_x", "X", tipo_campo, S.G_VARIABLES, True, True, "ayuda")
    assert c.opciones == opciones and c.editable and not c.obligatorio and c.clave is None   # no cuenta para la completitud


def test_desde_filas_ordena_y_descarta_lo_invalido():
    filas = [
        {"clave": "var_b", "etiqueta": "B", "tipo": "texto", "opciones": "", "activa": "False", "orden": "1"},
        {"clave": "var_a", "etiqueta": "A", "tipo": "opcion", "opciones": "x; y", "activa": "True", "orden": "0"},
        {"clave": "otra", "etiqueta": "mala", "tipo": "texto"},                              # sin prefijo
        {"clave": "var_c", "etiqueta": "C", "tipo": "inventado"},                           # tipo desconocido
        {"clave": "var_d", "etiqueta": "D", "tipo": "numero", "orden": "basura"},
    ]
    vs = V.desde_filas(filas)
    assert [v.clave for v in vs] == ["var_a", "var_d", "var_b"] or [v.clave for v in vs][:1] == ["var_a"]
    assert vs[0].opciones == ("x", "y") and not [v for v in vs if v.clave == "var_b"][0].activa


def test_canonizar_y_valores_invalidos():
    c = V.a_campo(V.Variable("var_t", "Tamaño", "opcion", TAMANO))
    assert V.canonizar(c, "grande") == "Grande" and V.canonizar(c, "pequeno") == "Pequeño" and V.canonizar(c, "Enorme") == "Enorme"
    assert V.valores_invalidos(c, "Grande") == [] and V.valores_invalidos(c, "Enorme") == ["Enorme"] and V.valores_invalidos(c, "") == []
    m = V.a_campo(V.Variable("var_m", "Varios", "etiquetas", TAMANO))
    assert V.canonizar(m, "grande; PEQUEÑO;otra") == "Grande; Pequeño; otra" and V.valores_invalidos(m, "Grande; otra") == ["otra"]
    libre = V.a_campo(V.Variable("var_l", "Libre", "texto"))
    assert V.canonizar(libre, "lo que sea") == "lo que sea" and V.valores_invalidos(libre, "lo que sea") == []
    assert V.valores_invalidos(S.CAMPO["fuente"], "otro.com") == []                              # solo las variables propias son cerradas


def test_validar_nueva_rechaza_nombres_repetidos():
    existentes = [V.Variable("var_tamano_del_proyecto", "Tamaño del proyecto", "opcion", TAMANO)]
    assert V.validar_nueva("Presupuesto", "numero", [], existentes) == ("Presupuesto", "var_presupuesto", ())
    for nombre in ("tamano del PROYECTO", "Título", "carpeta del proyecto", "  "):
        with pytest.raises(V.ErrorVariable):
            V.validar_nueva(nombre, "texto", [], existentes)
    with pytest.raises(V.ErrorVariable, match="tipo"):
        V.validar_nueva("Algo", "raro", [], existentes)
    with pytest.raises(V.ErrorVariable, match="nombre"):
        V.validar_nueva("¿¿??", "texto", [], existentes)
    with pytest.raises(V.ErrorVariable, match="demasiado largo"):
        V.validar_nueva("x" * 70, "texto", [], existentes)


# ============================================================================ registro dinámico
def test_registro_dinamico_en_el_esquema():
    base = len(S.COLUMNAS)
    assert S.registrar_variables([V.a_campo(V.Variable("var_t", "Tamaño", "opcion", TAMANO))]) is True
    assert len(S.COLUMNAS) == base + 1 and S.COLUMNAS[-1] == "var_t" and S.CAMPO["var_t"].variable
    assert S.etiqueta("var_t") == "Tamaño" and [c.key for c in S.variables_activas()] == ["var_t"]
    assert "var_t" in S.VISTAS["Proyecto y variables"] and S.VISTAS["Todo"][-1] == "var_t" and "var_t" not in S.VISTAS["Esenciales"]
    assert S.opciones_de("var_t") == list(TAMANO)
    assert S.registrar_variables([V.a_campo(V.Variable("var_t", "Tamaño", "opcion", TAMANO))]) is False   # idempotente
    S.registrar_variables([V.a_campo(V.Variable("var_t", "Tamaño", "opcion", TAMANO, activa=False))])
    assert S.variables_activas() == [] and "var_t" in S.CAMPO and "var_t" not in S.VISTAS["Proyecto y variables"]
    S.registrar_variables(())
    assert "var_t" not in S.CAMPO and len(S.COLUMNAS) == base and S.etiqueta("var_t") == "var_t"


def test_los_campos_base_no_cambian():
    assert S.COLUMNAS_SISTEMA == ("wp_id", "origen", "cargado_por", "fecha_carga", "editado_por", "fecha_edicion")
    assert S.CAMPO["carpeta_proyecto"].tipo == S.URL and S.CAMPO["documentos_proyecto"].tipo == S.LISTA_URL
    assert S.CAMPO["carpeta_proyecto"].grupo == S.G_PROYECTO and S.CAMPO["carpeta_proyecto"].editable
    assert len(S.CLAVE_CONTENIDO_CAMPOS) == 6 and len(S.CLAVE_ANALISIS_CAMPOS) == 8              # la completitud no cambia


# ============================================================================ crear (el ejemplo del pedido)
def test_crear_variable_tamano_del_proyecto(almacen):
    v = _crear(almacen, "tamaño_proyecto", desc="Pequeño: hasta 50 personas")
    assert (v.clave, v.etiqueta, v.tipo, v.opciones) == ("var_tamano_proyecto", "Tamaño proyecto", "opcion", TAMANO)
    # hoja Variables + columna nueva (vacía) en Base_Datos
    hoja = almacen.leer_hoja_texto(S.HOJA_VARIABLES)
    assert hoja.iloc[0]["clave"] == "var_tamano_proyecto" and hoja.iloc[0]["opciones"] == "Pequeño; Mediano; Grande"
    assert hoja.iloc[0]["creada_por"] == "Francis" and hoja.iloc[0]["activa"] == "True"
    base = almacen.leer_base_texto()
    assert "var_tamano_proyecto" in base.columns and (base["var_tamano_proyecto"] == "").all() and len(base) == 274
    # registrada en el esquema
    assert S.CAMPO["var_tamano_proyecto"].opciones == TAMANO and S.COLUMNAS[-1] == "var_tamano_proyecto"
    # libro de códigos actualizado con su definición
    lc = almacen.leer_hoja_texto(S.HOJA_LIBRO).set_index("columna")
    assert lc.at["var_tamano_proyecto", "opciones_respuesta"] == "Pequeño; Mediano; Grande"
    assert "Pequeño: hasta 50" in lc.at["var_tamano_proyecto", "descripcion"] and "Variable propia" in lc.at["var_tamano_proyecto", "fuente"]
    assert lc.at["var_tamano_proyecto", "tipo_variable"] == "Categórica única"


def test_las_variables_sobreviven_a_un_reinicio(almacen):
    _crear(almacen)
    S.registrar_variables(())                                    # como un proceso nuevo, sin nada registrado
    otro = LocalStorage(almacen.ruta)
    assert "var_tamano_del_proyecto" in otro.leer_base_texto().columns
    assert S.CAMPO["var_tamano_del_proyecto"].opciones == TAMANO
    S.registrar_variables(())
    with otro.transaccion():
        assert "var_tamano_del_proyecto" in S.CAMPO                 # también al abrir una transacción


def test_validaciones_al_crear(almacen):
    _crear(almacen)
    for kw, texto in [
        (dict(etiqueta="tamaño del proyecto"), "Ya existe"), (dict(etiqueta="Título"), "Ya existe"),
        (dict(etiqueta="Otra", opciones=["solo una"]), "al menos 2"), (dict(etiqueta="Otra", tipo="raro"), "tipo"),
        (dict(etiqueta=""), "nombre"), (dict(etiqueta="Otra", opciones=["a", "A", "b"]), "repetida"),
    ]:
        with pytest.raises(ErrorOperacion, match=texto):
            _crear(almacen, **{**dict(etiqueta="Otra", tipo="opcion", opciones=TAMANO), **kw})
    with pytest.raises(ErrorOperacion, match="nombre de quien edita"):
        _crear(almacen, "Nueva", usuario=" ")
    assert [v.clave for v in repo.listar_variables(almacen)] == ["var_tamano_del_proyecto"]


@pytest.mark.parametrize("tipo,opciones", [("texto", []), ("texto_largo", []), ("numero", []), ("fecha", []), ("url", []), ("si_no", [])])
def test_tipos_sin_lista_de_opciones(almacen, tipo, opciones):
    v = _crear(almacen, f"Var {tipo}", tipo, opciones)
    assert v.opciones == () and S.CAMPO[v.clave].variable


# ============================================================================ valores por noticia
def test_asignar_valores_con_historial(almacen):
    v = _crear(almacen)
    r = repo.aplicar_cambios([Cambio("EV0001", v.clave, "", "grande")], "Ana", almacen=almacen)
    assert r.ok and _valor(almacen, "EV0001", v.clave) == "Grande"                      # «grande» se guarda como la opción oficial
    h = repo.historial_df(almacen, "EV0001").iloc[0]
    assert (h["campo"], h["valor_antes"], h["valor_despues"], h["usuario"]) == (v.clave, "", "Grande", "Ana")
    assert _valor(almacen, "EV0001", "editado_por") == "Ana"


def test_valor_fuera_de_las_opciones_se_rechaza_con_mensaje(almacen):
    v = _crear(almacen)
    r = repo.aplicar_cambios([Cambio("EV0001", v.clave, "", "Enorme")], "Ana", almacen=almacen)
    assert not r.ok and "no es una opción de «Tamaño del proyecto»" in r.omitidos[0]["motivo"] and "Pequeño, Mediano, Grande" in r.omitidos[0]["motivo"]
    assert _valor(almacen, "EV0001", v.clave) == ""


def test_opciones_multiples_numero_y_si_no(almacen):
    multi = _crear(almacen, "Enfoques", "etiquetas", ["Agua", "Energía", "Suelo"])
    num = _crear(almacen, "Presupuesto", "numero", [])
    sn = _crear(almacen, "Tiene convenio", "si_no", [])
    r = repo.aplicar_cambios([
        Cambio("EV0001", multi.clave, "", "agua;SUELO"), Cambio("EV0001", num.clave, "", "1.234,5"),
        Cambio("EV0001", sn.clave, "", "si"), Cambio("EV0002", num.clave, "", "mucho"),
        Cambio("EV0002", multi.clave, "", "Agua; Fuego")], "Ana", almacen=almacen)
    assert _valor(almacen, "EV0001", multi.clave) == "Agua; Suelo"
    assert _valor(almacen, "EV0001", num.clave) == "1234.5"
    assert _valor(almacen, "EV0001", sn.clave) == "Sí"
    assert _valor(almacen, "EV0002", num.clave) == "mucho" and any("no es un número" in a for a in r.avisos)    # número: aviso, no bloqueo
    assert _valor(almacen, "EV0002", multi.clave) == "" and len(r.omitidos) == 1                              # lista cerrada: se rechaza


def test_agregar_filas_con_valor_invalido_lo_deja_vacio_y_avisa(almacen):
    v = _crear(almacen)
    r = repo.agregar_filas([{"titulo": "N", "url_noticia": "https://glocalminds.com/n/", "fuente": "glocalminds.com",
                             v.clave: "Gigante"},
                            {"titulo": "M", "url_noticia": "https://glocalminds.com/m/", "fuente": "glocalminds.com", v.clave: "mediano"}],
                           "Ana", almacen=almacen)
    assert r.ids_creados == ["EV0275", "EV0276"]
    assert _valor(almacen, "EV0275", v.clave) == "" and _valor(almacen, "EV0276", v.clave) == "Mediano"
    assert any("no es una opción válida; quedó vacío" in a for a in r.avisos)


def test_asignar_en_lote_opcion_unica_y_deshacer(almacen):
    v = _crear(almacen)
    r = repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002", "EV0003"], "Ana", "agregar", almacen)
    assert len(r.aplicados) == 3 and {_valor(almacen, i, v.clave) for i in ("EV0001", "EV0002")} == {"Grande"}
    repo.asignar_categoria(v.clave, "Pequeño", ["EV0001"], "Ana", "agregar", almacen)               # en opción única reemplaza
    assert _valor(almacen, "EV0001", v.clave) == "Pequeño"
    repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002"], "Ana", "quitar", almacen)       # quita solo si coincide
    assert _valor(almacen, "EV0001", v.clave) == "Pequeño" and _valor(almacen, "EV0002", v.clave) == ""
    lote = repo.lotes_df(almacen).iloc[0]
    repo.revertir_lote(lote["lote_id"], "Ana", almacen)
    assert _valor(almacen, "EV0002", v.clave) == "Grande"


def test_asignar_en_lote_opciones_multiples(almacen):
    v = _crear(almacen, "Enfoques", "etiquetas", ["Agua", "Energía", "Suelo"])
    repo.asignar_categoria(v.clave, "Agua", ["EV0001"], "Ana", "agregar", almacen)
    repo.asignar_categoria(v.clave, "Suelo", ["EV0001"], "Ana", "agregar", almacen)
    assert _valor(almacen, "EV0001", v.clave) == "Agua; Suelo"
    repo.asignar_categoria(v.clave, "Agua", ["EV0001"], "Ana", "quitar", almacen)
    assert _valor(almacen, "EV0001", v.clave) == "Suelo"
    with pytest.raises(ErrorOperacion, match="Dimensión no válida"):
        repo.asignar_categoria("titulo", "x", ["EV0001"], "Ana", "agregar", almacen)


# ============================================================================ gestionar opciones y estado
def test_uso_de_la_variable(almacen):
    v = _crear(almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002"], "Ana", "agregar", almacen)
    repo.asignar_categoria(v.clave, "Pequeño", ["EV0003"], "Ana", "agregar", almacen)
    u = repo.uso_variables(almacen)[v.clave]
    assert u["con_dato"] == 3 and u["por_opcion"] == {"Pequeño": 1, "Mediano": 0, "Grande": 2}


def test_agregar_y_quitar_opciones(almacen):
    v = _crear(almacen)
    repo.agregar_opcion(v.clave, "Muy grande", almacen)
    assert S.CAMPO[v.clave].opciones == TAMANO + ("Muy grande",)
    with pytest.raises(ErrorOperacion, match="repetida"):
        repo.agregar_opcion(v.clave, "grande", almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001"], "Ana", "agregar", almacen)
    with pytest.raises(ErrorOperacion, match="en uso en 1 noticia"):
        repo.quitar_opcion(v.clave, "Grande", almacen)                               # no se puede quitar lo que se usa
    repo.quitar_opcion(v.clave, "Muy grande", almacen)
    assert S.CAMPO[v.clave].opciones == TAMANO
    repo.quitar_opcion(v.clave, "Mediano", almacen)
    with pytest.raises(ErrorOperacion, match="al menos 2"):
        repo.quitar_opcion(v.clave, "Pequeño", almacen)
    with pytest.raises(ErrorOperacion, match="no existe"):
        repo.quitar_opcion(v.clave, "Fantasma", almacen)
    texto = _crear(almacen, "Libre", "texto", [])
    with pytest.raises(ErrorOperacion, match="no tiene lista de opciones"):
        repo.agregar_opcion(texto.clave, "x", almacen)


def test_renombrar_opcion_actualiza_las_noticias_en_un_lote(almacen):
    v = _crear(almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002"], "Ana", "agregar", almacen)
    r = repo.renombrar_opcion(v.clave, "Grande", "Extra grande", "Beto", almacen)
    assert len(r.aplicados) == 2 and S.CAMPO[v.clave].opciones == ("Pequeño", "Mediano", "Extra grande")
    assert _valor(almacen, "EV0001", v.clave) == "Extra grande"
    assert almacen.listar_respaldos(), "se guarda un respaldo antes de renombrar"
    with pytest.raises(ErrorOperacion, match="repetida"):
        repo.renombrar_opcion(v.clave, "Mediano", "pequeno", "Beto", almacen)
    repo.revertir_lote(r.lote_id, "Beto", almacen)                                    # los valores vuelven (la definición no)
    assert _valor(almacen, "EV0001", v.clave) == "Grande"


def test_actualizar_y_desactivar(almacen):
    v = _crear(almacen)
    repo.actualizar_variable(v.clave, almacen, descripcion="  nueva   descripción ")
    assert repo.listar_variables(almacen)[0].descripcion == "nueva descripción"
    repo.actualizar_variable(v.clave, almacen, etiqueta="Escala del proyecto")
    assert S.etiqueta(v.clave) == "Escala del proyecto" and v.clave == repo.listar_variables(almacen)[0].clave   # la columna no cambia
    with pytest.raises(ErrorOperacion, match="Ya existe"):
        repo.actualizar_variable(v.clave, almacen, etiqueta="Título")
    repo.asignar_categoria(v.clave, "Grande", ["EV0001"], "Ana", "agregar", almacen)
    repo.actualizar_variable(v.clave, almacen, activa=False)
    assert S.variables_activas() == [] and _valor(almacen, "EV0001", v.clave) == "Grande"                      # el dato se conserva
    assert almacen.leer_hoja_texto(S.HOJA_LIBRO).set_index("columna").at[v.clave, "descripcion"] == "nueva descripción"
    repo.actualizar_variable(v.clave, almacen, activa=True)
    assert [c.key for c in S.variables_activas()] == [v.clave]
    with pytest.raises(ErrorOperacion, match="no existe"):
        repo.actualizar_variable("var_fantasma", almacen, activa=False)


def test_la_variable_aparece_en_la_exportacion_y_la_tabla(almacen, monkeypatch):
    from tests.apptest_util import apuntar_a
    from utils import data, edicion, export
    v = _crear(almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001"], "Ana", "agregar", almacen)
    apuntar_a(monkeypatch, almacen)
    lectura = data.load_noticias()
    assert v.clave in lectura.columns and lectura.set_index("id_evento").at["EV0001", v.clave] == "Grande"
    assert (lectura[v.clave].isin(["", "Grande"])).all()                                       # sin NaN en las noticias sin valor
    assert ("var_tamano_del_proyecto", "Tamaño del proyecto", S.G_VARIABLES) in export.catalogo_exportable(lectura.columns)
    import io
    hojas = openpyxl.load_workbook(io.BytesIO(export.experiences_to_excel(lectura.head(2), [v.clave, "titulo"])))
    assert hojas["Noticias"]["A1"].value == "Tamaño del proyecto" and hojas["Noticias"]["A2"].value == "Grande"
    df = edicion.preparar_df(almacen.leer_base_texto(), data.opciones_de_campos())
    assert v.clave in df.columns and edicion.column_config(data.opciones_de_campos())[v.clave] is not None


# ============================================================================ migración de un libro anterior
def test_un_libro_anterior_recibe_las_columnas_nuevas_sin_perder_datos(almacen):
    wb = openpyxl.load_workbook(almacen.ruta)
    ws = wb[S.HOJA_BASE]
    cab = [c.value for c in ws[1]]
    for nombre in ("documentos_proyecto", "carpeta_proyecto"):                       # como un libro creado antes de este cambio
        ws.delete_cols(cab.index(nombre) + 1)
        cab.remove(nombre)
    del wb[S.HOJA_VARIABLES]
    wb.save(almacen.ruta)
    antes = {r: v for r, v in almacen.leer_base_texto().set_index("id_evento")["titulo"].items()}

    viejo = LocalStorage(almacen.ruta)
    viejo.asegurar()
    wb = openpyxl.load_workbook(almacen.ruta, read_only=True)
    try:
        assert S.HOJA_VARIABLES in wb.sheetnames
        assert [c for c in next(wb[S.HOJA_BASE].iter_rows(max_row=1, values_only=True)) if c][-2:] == ["carpeta_proyecto", "documentos_proyecto"]
    finally:
        wb.close()
    base = viejo.leer_base_texto()
    assert {r: v for r, v in base.set_index("id_evento")["titulo"].items()} == antes and (base["carpeta_proyecto"] == "").all()
    assert (base["origen"] == "historico").all()                                       # lo ya migrado no se vuelve a tocar
    firma = viejo.firma()
    viejo.asegurar()
    assert viejo.firma() == firma                                                      # idempotente
