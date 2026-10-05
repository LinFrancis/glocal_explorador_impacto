# -*- coding: utf-8 -*-
"""Variables propias y carpeta del proyecto, vistos desde las páginas: Administración, ficha, carga, tabla,
Explorador (filtros y análisis), exportación, importación y calidad."""
import io

import openpyxl
import pytest

from tests.apptest_util import apuntar_a, iniciar_sesion, nueva_app
from utils import calidad, export, importar, repo
from utils import schema as S
from utils.filters import CRITERIOS_KEY

TAMANO = ("Pequeño", "Mediano", "Grande")
CLAVE = "var_tamano_del_proyecto"


@pytest.fixture
def app(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    return at


def _crear(almacen, etiqueta="Tamaño del proyecto", tipo="opcion", opciones=TAMANO):
    return repo.crear_variable(etiqueta, tipo, list(opciones), "", "Ana", almacen)


def _fila(almacen, ide):
    return almacen.leer_base_texto().set_index("id_evento", drop=False).loc[ide]


def _ir(at, pagina):
    at.switch_page(pagina).run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def _boton(at, etiqueta, indice=0):
    return [b for b in at.button if b.label == etiqueta][indice]


def _tablas(at, columna):
    return [t.value for t in at.dataframe if columna in getattr(t.value, "columns", [])]


# ============================================================================ Administración: crear
def test_crear_variable_desde_administracion_el_ejemplo_del_pedido(app, almacen):
    _ir(app, "app_pages/administracion.py")
    app.text_input(key="var_nombre_0").set_value("tamaño_proyecto")
    app.text_area(key="var_opciones_0").set_value("pequeño\nmediana\ngrande")
    app.text_area(key="var_desc_0").set_value("Según el número de personas participantes").run()
    assert any("var_tamano_proyecto" in c.value for c in app.caption)              # avisa qué columna se creará
    _boton(app, "Crear variable").click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    v = repo.listar_variables(almacen)[0]
    assert (v.clave, v.etiqueta, v.tipo, v.opciones, v.descripcion) == (
        "var_tamano_proyecto", "Tamaño proyecto", "opcion", ("pequeño", "mediana", "grande"), "Según el número de personas participantes")
    assert any("creada" in s.value for s in app.success)
    assert app.session_state["var_form_n"] == 1 and app.text_input(key="var_nombre_1").value == ""    # formulario nuevo y vacío
    t = _tablas(app, "Variable")[0]
    assert t.iloc[0]["Variable"] == "Tamaño proyecto" and t.iloc[0]["Opciones"] == "pequeño, mediana, grande"


def test_el_tipo_decide_si_se_piden_opciones(app):
    _ir(app, "app_pages/administracion.py")
    assert [a.label for a in app.text_area if a.label.startswith("Opciones")]                 # opción única: pide opciones
    app.selectbox(key="var_tipo_0").select("numero").run()
    assert not [a for a in app.text_area if a.label.startswith("Opciones")]                    # número: no hay opciones
    app.selectbox(key="var_tipo_0").select("si_no").run()
    assert any("«Sí» y «No»" in c.value for c in app.caption)


def test_errores_de_creacion_se_explican(app, almacen):
    _ir(app, "app_pages/administracion.py")
    app.text_input(key="var_nombre_0").set_value("Mal")
    app.text_area(key="var_opciones_0").set_value("solo una")
    _boton(app, "Crear variable").click()
    app.run()
    assert any("al menos 2 opciones" in e.value for e in app.error) and repo.listar_variables(almacen) == []
    app.text_area(key="var_opciones_0").set_value("a\nb")
    app.text_input(key="var_nombre_0").set_value("Título")
    _boton(app, "Crear variable").click()
    app.run()
    assert any("Ya existe" in e.value for e in app.error)


def test_sin_nombre_de_editor_no_se_puede_crear(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at)
    at.session_state["auth_nombre"] = None
    _ir(at, "app_pages/administracion.py")
    assert _boton(at, "Crear variable").disabled


# ============================================================================ Administración: gestionar
def _seleccionar_variable(at):
    at.session_state["var_tabla"] = {"selection": {"rows": [0], "columns": [], "cells": []}}


def test_gestionar_opciones_agregar_y_quitar(app, almacen):
    v = _crear(almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001"], "Ana", "agregar", almacen)
    _ir(app, "app_pages/administracion.py")
    _seleccionar_variable(app)
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    ops = _tablas(app, "Opción")[0]
    assert ops["Opción"].tolist() == list(TAMANO) and ops["Noticias"].tolist() == [0, 0, 1]      # con el uso de cada una
    app.text_input(key=f"var_nueva_op_{v.clave}_3").set_value("Muy grande")
    _seleccionar_variable(app)
    app.run()                                                                # ahora el botón «Agregar» ya está activo
    _seleccionar_variable(app)
    _boton(app, "Agregar").click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    assert S.CAMPO[v.clave].opciones == TAMANO + ("Muy grande",)
    # quitar una opción en uso: se rechaza con mensaje
    app.selectbox(key=f"var_op_sel_{v.clave}").select("Grande")
    _seleccionar_variable(app)
    app.run()
    _seleccionar_variable(app)
    _boton(app, "Quitar opción").click()
    app.run()
    assert any("en uso en 1 noticia" in e.value for e in app.error) and "Grande" in S.CAMPO[v.clave].opciones


def test_renombrar_y_desactivar_estan_protegidos_con_clave(app, almacen):
    v = _crear(almacen)
    _ir(app, "app_pages/administracion.py")
    _seleccionar_variable(app)
    app.run()
    assert any(b.label == "Desactivar variable" for b in app.button) and any(b.label == "Renombrar opción" for b in app.button)
    _seleccionar_variable(app)
    _boton(app, "Desactivar variable").click()
    app.run()
    assert not app.exception and S.variables_activas()                      # abrir el diálogo no desactiva nada


# ============================================================================ Administración: asignar en lote
def test_asignar_la_variable_en_lote_desde_administracion(app, almacen):
    v = _crear(almacen)
    _ir(app, "app_pages/administracion.py")
    assert "Tamaño del proyecto" in app.selectbox(key="asig_dim").options
    app.selectbox(key="asig_dim").select(v.clave).run()
    assert [r.label for r in app.radio if r.key and r.key.startswith("asig_accion")][0] == "Acción"
    assert app.radio(key="asig_accion_u").options == ["Asignar", "Quitar"]                      # opción única: «Asignar»
    app.selectbox(key=f"asig_cat_{v.clave}").select("Mediano").run()
    app.text_input(key="asig_texto").set_value("EV0001").run()
    app.checkbox(key=f"asig_todas_{v.clave}_Asignar_Mediano").check().run()
    boton = next(b for b in app.button if b.label.startswith("Asignar «Mediano»"))
    assert boton.label == "Asignar «Mediano» en 1 noticia(s)"
    boton.click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    assert _fila(almacen, "EV0001")[v.clave] == "Mediano"
    assert repo.lotes_df(almacen).iloc[0]["accion"] == "Categoría"


# ============================================================================ Ficha: variables y carpeta del proyecto
def _guardar(at, indice):
    [b for b in at.button if b.label == "Guardar cambios"][indice].click()
    at.run()
    assert not at.exception, [e.value for e in at.exception]


def test_ficha_asigna_la_variable_y_guarda_con_historial(app, almacen):
    v = _crear(almacen)
    _ir(app, "app_pages/ficha.py")
    sel = app.selectbox(key=f"ficha_EV0001_0_{v.clave}")
    assert sel.options == list(TAMANO) and sel.value is None                                     # sin valor al empezar
    sel.select("Grande")
    _guardar(app, 2)                                                                              # 3.er formulario: variables propias
    assert _fila(almacen, "EV0001")[v.clave] == "Grande"
    h = repo.historial_df(almacen, "EV0001").iloc[0]
    assert (h["campo"], h["valor_despues"], h["usuario"]) == (v.clave, "Grande", "Francis")


def test_ficha_sin_variables_invita_a_crearlas(app):
    _ir(app, "app_pages/ficha.py")
    assert any("Aún no hay variables propias" in c.value for c in app.caption)


def test_ficha_carpeta_del_proyecto_y_documentos(app, almacen):
    _ir(app, "app_pages/ficha.py")
    assert not [b for b in app.get("link_button") if "carpeta del proyecto" in b.proto.label]
    app.text_input(key="ficha_EV0001_0_carpeta_proyecto").set_value("drive.google.com/drive/folders/abc123")
    app.text_area(key="ficha_EV0001_0_documentos_proyecto").set_value("https://docs.google.com/d/1\nhttps://x.cl/informe.pdf")
    _guardar(app, 0)
    f = _fila(almacen, "EV0001")
    assert f["carpeta_proyecto"] == "https://drive.google.com/drive/folders/abc123"                # se le agrega https://
    assert f["documentos_proyecto"] == "https://docs.google.com/d/1 | https://x.cl/informe.pdf"
    botones = [b.proto.label for b in app.get("link_button")]
    assert "Abrir carpeta del proyecto" in botones                                                  # botón destacado en la cabecera
    assert set(repo.historial_df(almacen, "EV0001")["campo"]) == {"carpeta_proyecto", "documentos_proyecto"}


def test_ficha_vista_de_lectura_muestra_carpeta_y_variables(app, almacen):
    v = _crear(almacen)
    repo.aplicar_cambios([repo.Cambio("EV0001", v.clave, "", "Grande"),
                          repo.Cambio("EV0001", "carpeta_proyecto", "", "https://drive.google.com/x"),
                          repo.Cambio("EV0001", "documentos_proyecto", "", "https://a.cl/1 | https://a.cl/2")], "Ana", almacen=almacen)
    _ir(app, "app_pages/ficha.py")
    assert any("Tamaño del proyecto" in m.value for m in app.markdown)
    assert any("Grande" in m.value for m in app.markdown)
    assert "Abrir carpeta del proyecto ↗" in [b.proto.label for b in app.get("link_button")]
    assert any("Otros enlaces del proyecto (2)" in e.label for e in app.expander)


# ============================================================================ Carga
def test_carga_formulario_incluye_carpeta_y_variables(app, almacen):
    v = _crear(almacen)
    _ir(app, "app_pages/carga.py")
    app.text_input(key="carga_0_titulo").set_value("Proyecto con variables")
    app.text_input(key="carga_0_url_noticia").set_value("https://glocalminds.com/catalogo/pv/")
    app.selectbox(key="carga_0_fuente").select("glocalminds.com")
    app.text_input(key="carga_0_carpeta_proyecto").set_value("https://drive.google.com/drive/folders/zzz")
    app.selectbox(key=f"carga_0_{v.clave}").select("Mediano")
    _boton(app, "Agregar noticia").click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    f = _fila(almacen, "EV0275")
    assert (f["carpeta_proyecto"], f[v.clave], f["titulo"]) == ("https://drive.google.com/drive/folders/zzz", "Mediano", "Proyecto con variables")


# ============================================================================ Tabla de datos
def test_tabla_vista_proyecto_y_variables(app, almacen):
    v = _crear(almacen)
    _ir(app, "app_pages/tabla.py")
    assert "Proyecto y variables" in app.segmented_control(key="tabla_vista").options
    app.segmented_control(key="tabla_vista").set_value("Proyecto y variables").run()
    assert not app.exception, [e.value for e in app.exception]
    # una edición válida se guarda...
    valida = {"edited_rows": {0: {v.clave: "Grande"}}, "added_rows": [], "deleted_rows": []}
    app.session_state["tabla_editor_0"] = valida
    app.run()
    app.session_state["tabla_editor_0"] = valida
    _boton(app, "Guardar 1 cambio(s)").click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    assert _fila(almacen, "EV0001")[v.clave] == "Grande"
    # ...y una que no es una opción de la variable no se puede guardar (el libro de códigos la rechaza)
    app.session_state["tabla_editor_1"] = {"edited_rows": {1: {v.clave: "Enorme"}}, "added_rows": [], "deleted_rows": []}
    app.run()
    assert any("no respetan el libro de códigos" in e.value for e in app.error)
    assert any("«Tamaño del proyecto»: «Enorme» no es una de sus opciones" in c.value for c in app.caption)
    assert next(b for b in app.button if b.label.startswith("Guardar")).disabled
    assert _fila(almacen, "EV0002")[v.clave] == ""


# ============================================================================ Explorador: filtros y análisis
@pytest.fixture
def con_datos(almacen):
    v = _crear(almacen)
    num = _crear(almacen, "Presupuesto", "numero", [])
    multi = _crear(almacen, "Enfoques", "etiquetas", ["Agua", "Energía"])
    grandes = [f"EV{i:04d}" for i in range(1, 6)]
    repo.asignar_categoria(v.clave, "Grande", grandes, "Ana", "agregar", almacen)
    repo.asignar_categoria(v.clave, "Pequeño", ["EV0006", "EV0007", "EV0008"], "Ana", "agregar", almacen)
    repo.asignar_categoria(multi.clave, "Agua", ["EV0001", "EV0006"], "Ana", "agregar", almacen)
    repo.aplicar_cambios([repo.Cambio(f"EV{i:04d}", num.clave, "", str(1000 * i)) for i in range(1, 6)], "Ana", almacen=almacen)
    return v, num, multi


def _encontradas(at):
    return next(m.value for m in at.metric if m.label == "Experiencias encontradas")


def test_filtrar_por_variable_propia(app, con_datos):
    v, num, multi = con_datos
    _ir(app, "app_pages/explorador.py")
    assert app.multiselect(key=f"f_var_{v.clave}").options == list(TAMANO) + ["(sin dato)"]
    app.multiselect(key=f"f_var_{v.clave}").set_value(["Grande"]).run()
    assert _encontradas(app) == "5 de 274"
    app.multiselect(key=f"f_var_{v.clave}").set_value(["Grande", "Pequeño"]).run()
    assert _encontradas(app) == "8 de 274"
    app.multiselect(key=f"f_var_{v.clave}").set_value(["(sin dato)"]).run()
    assert _encontradas(app) == "266 de 274"                                                      # las que no tienen valor
    assert app.session_state[CRITERIOS_KEY]["variables"] == {v.clave: ["(sin dato)"]}
    app.multiselect(key=f"f_var_{v.clave}").set_value([]).run()
    app.multiselect(key=f"f_var_{multi.clave}").set_value(["Agua"]).run()                          # opción múltiple
    assert _encontradas(app) == "2 de 274"


def test_filtro_numerico_y_limpiar(app, con_datos):
    v, num, _ = con_datos
    _ir(app, "app_pages/explorador.py")
    app.slider(key=f"f_var_{num.clave}").set_range(2000, 4000).run()
    assert _encontradas(app) == "3 de 274"
    _boton(app, "Limpiar filtros").click().run()
    assert _encontradas(app) == "274 de 274"
    assert not (app.session_state[CRITERIOS_KEY].get("variables") if CRITERIOS_KEY in app.session_state else None)


def test_el_filtro_de_variable_persiste_al_volver(app, con_datos):
    v, _, _ = con_datos
    _ir(app, "app_pages/explorador.py")
    app.multiselect(key=f"f_var_{v.clave}").set_value(["Grande"]).run()
    _ir(app, "app_pages/glosario.py")
    _ir(app, "app_pages/explorador.py")
    assert app.multiselect(key=f"f_var_{v.clave}").value == ["Grande"] and _encontradas(app) == "5 de 274"


def test_analisis_distribucion_y_cruce(app, con_datos):
    v, num, multi = con_datos
    _ir(app, "app_pages/explorador.py")
    app.selectbox(key="var_an_clave").select(v.clave).run()
    assert not app.exception, [e.value for e in app.exception]
    m = {x.label: x.value for x in app.metric}
    assert m["Noticias con dato"] == "8 de 274" and m["Cobertura"] == "3%"
    dist = _tablas(app, "% de las noticias")[0]
    assert dist["Opción"].tolist() == ["Pequeño", "Grande", "(sin dato)"] or dist["Opción"].tolist() == ["Pequeño", "Mediano", "Grande", "(sin dato)"][:len(dist)]
    assert dict(zip(dist["Opción"], dist["Noticias"])) == {"Pequeño": 3, "Grande": 5, "(sin dato)": 266}
    assert len(app.get("plotly_chart")) >= 1
    # cruce con otra dimensión
    app.selectbox(key="var_an_cruce").select("categoria_macro").run()
    assert not app.exception, [e.value for e in app.exception]
    assert any(k == "download_button" for k in [e.type for e in app.get("download_button")]) or app.get("download_button")
    # variable numérica: estadísticas
    app.selectbox(key="var_an_clave").select(num.clave).run()
    assert not app.exception
    app.selectbox(key="var_an_cruce").select("(ninguna)").run()
    m = {x.label: x.value for x in app.metric}
    assert (m["Mínimo"], m["Máximo"], m["Mediana"]) == ("1000", "5000", "3000")


def test_analisis_sin_variables_explica_como_crearlas(app):
    _ir(app, "app_pages/explorador.py")
    assert any("Todavía no hay variables propias" in i.value for i in app.info)


def test_las_variables_inactivas_no_se_ofrecen(app, almacen, con_datos):
    v, _, _ = con_datos
    repo.actualizar_variable(v.clave, almacen, activa=False)
    _ir(app, "app_pages/explorador.py")
    assert f"f_var_{v.clave}" not in [m.key for m in app.multiselect]
    assert v.clave not in app.selectbox(key="var_an_clave").options


# ============================================================================ Exportación
def test_la_exportacion_ofrece_las_variables_como_grupo(app, con_datos):
    v, _, _ = con_datos
    _ir(app, "app_pages/explorador.py")
    app.radio(key="exp_info").set_value("Solo las columnas que elijo").run()
    grupo = app.multiselect(key="exp_cols_Variables propias")
    assert "Tamaño del proyecto" in grupo.options
    grupo.set_value([v.clave]).run()
    assert not app.exception, [e.value for e in app.exception]
    assert any("columna(s) elegida(s)" in c.value for c in app.caption)
    assert "exp_cols_Carpeta y documentos" in [m.key for m in app.multiselect]


def test_excel_y_word_completos_incluyen_variables_y_carpeta(almacen, monkeypatch):
    from docx import Document

    from utils import data
    v = _crear(almacen)
    repo.aplicar_cambios([repo.Cambio("EV0001", v.clave, "", "Grande"), repo.Cambio("EV0001", "carpeta_proyecto", "", "https://drive.google.com/x")],
                         "Ana", almacen=almacen)
    apuntar_a(monkeypatch, almacen)
    df = data.load_noticias().head(1)
    hojas = openpyxl.load_workbook(io.BytesIO(export.experiences_to_excel(df)))
    resumen = [c.value for c in hojas["Resumen"][1]]
    assert "Tamaño del proyecto" in resumen and "Carpeta del proyecto" in resumen
    assert hojas["Resumen"].cell(2, resumen.index("Tamaño del proyecto") + 1).value == "Grande"
    doc = Document(io.BytesIO(export.experiences_to_word(df)))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Tamaño del proyecto: Grande" in texto and "Carpeta del proyecto: https://drive.google.com/x" in texto


# ============================================================================ Importación
def test_plantilla_y_mapeo_incluyen_carpeta_y_variables(almacen):
    v = _crear(almacen)
    num = _crear(almacen, "Presupuesto", "numero", [])
    datos = importar.plantilla_xlsx()
    wb = openpyxl.load_workbook(io.BytesIO(datos))
    encabezados = [c.value for c in wb["Noticias"][1]]
    assert "Carpeta del proyecto" in encabezados and "Tamaño del proyecto" in encabezados and "Presupuesto" in encabezados
    ayuda = {r[0].value: r[3].value for r in wb["Instrucciones"].iter_rows(min_row=2) if r[0].value}
    assert "Opciones: Pequeño; Mediano; Grande" in ayuda["Tamaño del proyecto"]
    mapa = importar.mapear_columnas(["Título", "Enlace", "Tamaño del proyecto", "presupuesto", "Carpeta del proyecto"])
    assert mapa["Tamaño del proyecto"] == v.clave and mapa["presupuesto"] == num.clave and mapa["Carpeta del proyecto"] == "carpeta_proyecto"


def test_importar_con_variables_normaliza_y_avisa(almacen):
    v = _crear(almacen)
    filas = importar.filas_desde_df(
        importar.leer_archivo("Título,Enlace,Tamaño del proyecto,Carpeta del proyecto\n"
                              "A,https://glocalminds.com/catalogo/a/,grande,drive.google.com/x\n"
                              "B,https://glocalminds.com/catalogo/b/,Gigante,\n".encode(), "d.csv"),
        importar.mapear_columnas(["Título", "Enlace", "Tamaño del proyecto", "Carpeta del proyecto"]))
    res = repo.agregar_filas(filas, "Ana", accion=S.ACC_CARGA, origen="archivo", almacen=almacen)
    assert res.ids_creados == ["EV0275", "EV0276"]
    assert _fila(almacen, "EV0275")[v.clave] == "Grande" and _fila(almacen, "EV0275")["carpeta_proyecto"] == "https://drive.google.com/x"
    assert _fila(almacen, "EV0276")[v.clave] == "" and any("no es una opción válida" in a for a in res.avisos)


# ============================================================================ Calidad
def test_calidad_detecta_valores_fuera_de_las_opciones_y_enlaces_malos(almacen):
    v = _crear(almacen)
    repo.aplicar_cambios([repo.Cambio("EV0001", "carpeta_proyecto", "", "esto no es un enlace"),
                          repo.Cambio("EV0002", "documentos_proyecto", "", "https://ok.cl/a | otra cosa rara")], "Ana", almacen=almacen)
    wb = openpyxl.load_workbook(almacen.ruta)                                   # alguien edita el Excel a mano
    ws = wb[S.HOJA_BASE]
    cab = [c.value for c in ws[1]]
    ws.cell(row=2, column=cab.index(v.clave) + 1).value = "Enorme"
    wb.save(almacen.ruta)
    base = almacen.leer_base_texto()
    fuera = calidad.variables_fuera_de_opciones(base)
    assert [(f["id_evento"], f["actual"]) for f in fuera.filas] == [("EV0001", "Enorme")] and "Pequeño" in fuera.filas[0]["sugerido"]
    assert fuera.columna_sugerida == "Opciones válidas" and not fuera.corregible
    urls = {(f["id_evento"], f["campo"]) for f in calidad.urls_mal_formadas(base).filas}
    assert urls == {("EV0001", "carpeta_proyecto"), ("EV0002", "documentos_proyecto")}
    assert any(h.clave == "variables_opciones" for h in calidad.detectar(base))
