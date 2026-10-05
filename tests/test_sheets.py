# -*- coding: utf-8 -*-
"""Almacén en Google Sheets, probado con una hoja simulada en memoria (sin red ni credenciales)."""
import io

import openpyxl
import pandas as pd
import pytest

from utils import repo, sheets
from utils import schema as S
from utils.repo import Cambio
from utils.sheets import ClienteMemoria, SheetsStorage
from utils.storage import ErrorAlmacen

USUARIO = "Ana"


@pytest.fixture(scope="session")
def hojas_sembradas(semilla):
    """Contenido de una hoja de Google recién llenada desde el Excel original (se hace una vez: tarda)."""
    cliente = ClienteMemoria()
    SheetsStorage(cliente, semilla, ttl=0).asegurar()
    return {k: [list(f) for f in v] for k, v in cliente.hojas.items()}


@pytest.fixture
def cliente(hojas_sembradas):
    return ClienteMemoria(hojas_sembradas)


@pytest.fixture
def hoja(cliente, semilla) -> SheetsStorage:
    alm = SheetsStorage(cliente, semilla, ttl=0)
    alm.asegurar()
    cliente.llamadas = {k: 0 for k in cliente.llamadas}
    cliente.filas_escritas = 0
    return alm


def _nueva(i=1, **extra):
    return {"titulo": f"Noticia {i}", "url_noticia": f"https://glocalminds.com/catalogo/n{i}/",
            "fuente": "glocalminds.com", "slug": f"n{i}", "wp_id": str(9000 + i), **extra}


def _celda(cliente, hoja, ide, campo):
    m = cliente.hojas[hoja]
    fila = next(f for f in m[1:] if f[m[0].index("id_evento")] == ide)
    return fila[m[0].index(campo)] if campo in m[0] else None


# ============================================================================ primera carga
def test_una_hoja_vacia_se_llena_desde_el_excel_original(semilla):
    cliente = ClienteMemoria({"Hoja 1": [[""]]})                 # una hoja nueva de Google trae una pestaña vacía
    alm = SheetsStorage(cliente, semilla, ttl=0)
    alm.asegurar()
    assert set(cliente.hojas) == {S.HOJA_BASE, S.HOJA_LIBRO, *S.HOJAS_AUX}        # y la pestaña por defecto se quita
    assert len(cliente.hojas[S.HOJA_BASE]) == 275                                # encabezado + 274 noticias
    assert set(S.COLUMNAS_NUEVAS) <= set(cliente.hojas[S.HOJA_BASE][0])
    assert cliente.hojas[S.HOJA_CATEGORIAS][1][0] in S.DIMENSIONES_CATALOGO      # catálogo sembrado
    assert len(alm.leer_base_texto()) == 274


def test_una_hoja_vacia_prefiere_el_libro_local_si_existe(tmp_path, plantilla, semilla):
    import shutil
    local = tmp_path / "local.xlsx"
    shutil.copy2(plantilla, local)
    cliente = ClienteMemoria()
    SheetsStorage(cliente, semilla, inicial=local, ttl=0).asegurar()
    assert len(cliente.hojas[S.HOJA_BASE]) == 275 and S.HOJA_EDITORES in cliente.hojas


def test_sin_semilla_ni_libro_se_explica(tmp_path):
    with pytest.raises(ErrorAlmacen, match="está vacía"):
        SheetsStorage(ClienteMemoria(), tmp_path / "no_existe.xlsx", ttl=0).asegurar()


def test_asegurar_es_idempotente_y_no_reescribe(hoja, cliente):
    hoja.asegurar()
    assert cliente.llamadas["reemplazar"] == 0 and cliente.llamadas["escribir_filas"] == 0


def test_lo_que_se_lee_es_igual_que_con_el_archivo_local(hoja, almacen):
    pd.testing.assert_frame_equal(hoja.leer_base_texto(), almacen.leer_base_texto())
    pd.testing.assert_frame_equal(hoja.leer_hoja_texto(S.HOJA_LIBRO), almacen.leer_hoja_texto(S.HOJA_LIBRO))
    cols = ["dimension", "nombre", "descripcion", "activa", "creada_por"]                  # la fecha de siembra difiere
    pd.testing.assert_frame_equal(hoja.leer_hoja_texto(S.HOJA_CATEGORIAS)[cols], almacen.leer_hoja_texto(S.HOJA_CATEGORIAS)[cols])
    local, remoto = almacen.leer_base_excel(), hoja.leer_base_excel()
    assert list(local.columns) == list(remoto.columns) and local.shape == remoto.shape
    assert str(local["fecha_publicacion_web"].dtype) == str(remoto["fecha_publicacion_web"].dtype)
    assert local["titulo"].tolist() == remoto["titulo"].tolist()


def test_exportar_xlsx_entrega_un_libro_completo(hoja):
    wb = openpyxl.load_workbook(io.BytesIO(hoja.exportar_xlsx()))
    assert S.HOJA_BASE in wb.sheetnames and wb[S.HOJA_BASE].max_row == 275


# ============================================================================ escritura
def test_un_cambio_sube_solo_la_fila_modificada(hoja, cliente):
    res = repo.aplicar_cambios([Cambio("EV0003", "autor", "", "Ana Soto")], USUARIO, almacen=hoja)
    assert res.ok and len(res.aplicados) == 1
    assert _celda(cliente, S.HOJA_BASE, "EV0003", "autor") == "Ana Soto"
    assert _celda(cliente, S.HOJA_BASE, "EV0003", "editado_por") == USUARIO
    assert cliente.llamadas["reemplazar"] == 0                      # no se reescribe la base completa
    assert cliente.filas_escritas <= 4                              # la noticia + la fila nueva del historial
    assert hoja.leer_base_texto().set_index("id_evento").at["EV0003", "autor"] == "Ana Soto"
    assert len(repo.historial_df(hoja)) == 1


def test_si_la_transaccion_falla_no_se_sube_nada(hoja, cliente):
    version = cliente.version
    with pytest.raises(RuntimeError):
        with hoja.transaccion() as libro:
            libro.set_valor("EV0001", "autor", "No debería quedar")
            raise RuntimeError("falla a mitad de camino")
    assert cliente.version == version and _celda(cliente, S.HOJA_BASE, "EV0001", "autor") == ""
    assert hoja.leer_base_texto().set_index("id_evento").at["EV0001", "autor"] == ""


def test_agregar_noticias_y_deshacer_el_lote(hoja, cliente):
    res = repo.agregar_filas([_nueva(1), _nueva(2)], USUARIO, accion=S.ACC_CARGA, origen="archivo", almacen=hoja)
    assert res.ids_creados == ["EV0275", "EV0276"] and len(cliente.hojas[S.HOJA_BASE]) == 277
    assert _celda(cliente, S.HOJA_BASE, "EV0276", "origen") == "archivo"
    assert cliente.llamadas["reemplazar"] == 0                      # solo filas nuevas al final
    r = repo.revertir_lote(res.lote_id, "Luis", almacen=hoja)
    assert r.ok and len(hoja.leer_base_texto()) == 274 and len(cliente.hojas[S.HOJA_BASE]) == 275


def test_los_tipos_viajan_como_texto_y_vuelven_tipados(hoja, cliente):
    repo.aplicar_cambios([Cambio("EV0001", "es_duplicado_secundario", "False", "True"),
                          Cambio("EV0001", "fecha_publicacion_web", hoja.leer_base_texto().at[0, "fecha_publicacion_web"], "2020-05-17")],
                         USUARIO, almacen=hoja)
    assert _celda(cliente, S.HOJA_BASE, "EV0001", "es_duplicado_secundario") == "True"
    assert _celda(cliente, S.HOJA_BASE, "EV0001", "fecha_publicacion_web") == "2020-05-17"
    df = hoja.leer_base_excel()
    assert bool(df.loc[0, "es_duplicado_secundario"]) is True and str(df.loc[0, "fecha_publicacion_web"]).startswith("2020-05-17")


def test_los_textos_muy_largos_se_recortan_al_limite_de_google(hoja, cliente):
    repo.aplicar_cambios([Cambio("EV0001", "contenido_completo", hoja.leer_base_texto().at[0, "contenido_completo"], "x" * 60000)],
                         USUARIO, almacen=hoja)
    assert len(_celda(cliente, S.HOJA_BASE, "EV0001", "contenido_completo")) == sheets.MAX_CELDA


def test_variables_propias_viven_en_la_hoja(hoja, cliente):
    v = repo.crear_variable("Tamaño del proyecto", "opcion", ["Pequeño", "Mediano", "Grande"], "", USUARIO, hoja)
    assert v.clave in cliente.hojas[S.HOJA_BASE][0]
    assert cliente.hojas[S.HOJA_VARIABLES][1][0] == v.clave
    repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002"], USUARIO, "agregar", hoja)
    assert _celda(cliente, S.HOJA_BASE, "EV0002", v.clave) == "Grande"
    libro = {f[0]: f for f in cliente.hojas[S.HOJA_LIBRO][1:]}
    assert libro[v.clave][3] == "Pequeño; Mediano; Grande"          # el libro de códigos la trae
    # otro proceso que abra la misma hoja la ve completa, con la variable ya registrada
    S.registrar_variables(())
    otro = SheetsStorage(cliente, ttl=0)
    assert otro.leer_base_texto().set_index("id_evento").at["EV0001", v.clave] == "Grande"
    assert v.clave in S.CAMPO


def test_el_equipo_se_registra_en_la_hoja(hoja, cliente):
    assert repo.registrar_editor("María José", hoja) == "María José"
    assert repo.registrar_editor("maria jose", hoja) == "María José"
    assert cliente.hojas[S.HOJA_EDITORES][1][0] == "María José" and len(cliente.hojas[S.HOJA_EDITORES]) == 2


# ============================================================================ cambios hechos por otras personas
def test_una_edicion_a_mano_en_google_se_ve_al_refrescar(hoja, cliente):
    f0 = hoja.firma()
    col = cliente.hojas[S.HOJA_BASE][0].index("autor") + 1
    cliente.editar_a_mano(S.HOJA_BASE, 2, col, "Escrito en Google Sheets")
    assert hoja.firma() != f0
    assert hoja.leer_base_texto().at[0, "autor"] == "Escrito en Google Sheets"


def test_conflicto_si_otra_persona_cambio_la_celda_entre_medio(hoja, cliente):
    col = cliente.hojas[S.HOJA_BASE][0].index("autor") + 1
    cliente.editar_a_mano(S.HOJA_BASE, 2, col, "Otra persona")
    res = repo.aplicar_cambios([Cambio("EV0001", "autor", "", "Yo")], USUARIO, almacen=hoja)      # yo creía que estaba vacío
    assert not res.aplicados and res.omitidos and _celda(cliente, S.HOJA_BASE, "EV0001", "autor") == "Otra persona"


def test_la_consulta_de_version_se_limita_con_el_ttl(cliente, semilla):
    alm = SheetsStorage(cliente, semilla, ttl=60)
    alm.asegurar()
    cliente.llamadas["marca"] = 0
    for _ in range(20):
        alm.firma()
        alm.leer_hoja_texto(S.HOJA_NOTAS)
    assert cliente.llamadas["marca"] == 0 and cliente.llamadas["leer"] == 1         # lecturas desde la copia en memoria


# ============================================================================ respaldos
def test_respaldos_en_pestanas_y_limite(hoja, cliente):
    for i in range(7):
        assert hoja.respaldar(f"t{i}")
    nombres = [t for t in cliente.hojas if t.startswith(sheets.PREFIJO_RESPALDO)]
    assert len(nombres) == 5 and len(hoja.respaldos()) == 5
    assert all(r.fecha is not None for r in hoja.respaldos())
    assert sheets.PREFIJO_RESPALDO not in "".join(hoja.leer_base_texto().columns)
    assert not any(t.startswith(sheets.PREFIJO_RESPALDO) for t in cliente.leer())     # no se descargan en cada lectura


def test_respaldo_antes_de_una_operacion_masiva(hoja):
    repo.agregar_filas([_nueva(1)], USUARIO, accion=S.ACC_CARGA, origen="archivo", almacen=hoja, respaldo="antes_de_importar")
    assert hoja.listar_respaldos() and hoja.listar_respaldos()[0].endswith("antes_de_importar")


def test_restablecer_desde_el_original(hoja, cliente):
    repo.aplicar_cambios([Cambio("EV0001", "autor", "", "X")], USUARIO, almacen=hoja)
    respaldo = hoja.restablecer_desde_semilla()
    assert respaldo and respaldo.endswith("antes_de_restablecer")
    assert _celda(cliente, S.HOJA_BASE, "EV0001", "autor") == "" and hoja.leer_hoja_texto(S.HOJA_HISTORIAL).empty
    assert len(hoja.leer_base_texto()) == 274


# ============================================================================ configuración
@pytest.fixture
def secretos(monkeypatch):
    monkeypatch.delenv("IMPACTO_ALMACEN", raising=False)
    datos = {}
    monkeypatch.setattr(sheets, "_seccion_de_secretos", lambda: datos or None)
    return datos


def test_sin_secretos_se_usa_el_archivo_local(secretos):
    assert sheets.configuracion() is None


def test_las_pruebas_fuerzan_el_archivo_local_aunque_haya_secretos(secretos, monkeypatch):
    secretos.update(spreadsheet="abc", client_email="a@b.c", private_key="k")
    monkeypatch.setenv("IMPACTO_ALMACEN", "local")
    assert sheets.configuracion() is None


def test_configuracion_completa(secretos):
    secretos.update(spreadsheet="https://docs.google.com/spreadsheets/d/ABC/edit", type="service_account", project_id="p",
                    client_email="plataforma@p.iam.gserviceaccount.com",
                    private_key="-----BEGIN PRIVATE KEY-----\\nAAA\\n-----END PRIVATE KEY-----\\n")
    cfg = sheets.configuracion()
    assert cfg["spreadsheet"].endswith("/ABC/edit")
    assert cfg["info"]["private_key"] == "-----BEGIN PRIVATE KEY-----\nAAA\n-----END PRIVATE KEY-----\n"     # \n literales -> saltos
    assert "spreadsheet" not in cfg["info"] and cfg["info"]["token_uri"].startswith("https://")


def test_la_direccion_puede_llamarse_de_otra_forma(secretos):
    secretos.update(spreadsheet_id="ABC", client_email="a@b.c", private_key="k")
    assert sheets.configuracion()["spreadsheet"] == "ABC"


@pytest.mark.parametrize("secreto,pista", [
    ({"client_email": "a@b.c", "private_key": "k"}, "spreadsheet"),
    ({"spreadsheet": "ABC", "client_email": "a@b.c"}, "private_key"),
    ({"spreadsheet": "ABC"}, "client_email"),
])
def test_configuracion_incompleta_avisa_en_vez_de_caer_al_archivo_local(secretos, secreto, pista):
    secretos.update(secreto)
    with pytest.raises(ErrorAlmacen, match="incompleta") as e:
        sheets.configuracion()
    assert pista in str(e.value)


def test_get_storage_elige_el_almacen_segun_los_secretos(secretos, monkeypatch, cliente, semilla):
    from utils import storage
    storage.olvidar_almacenes()
    secretos.update(spreadsheet="ABC", client_email="a@b.c", private_key="k")
    monkeypatch.setattr(sheets, "ClienteGspread", lambda info, hoja: cliente)
    alm = storage.get_storage()
    try:
        assert isinstance(alm, SheetsStorage) and alm.url.startswith("https://docs.google.com/spreadsheets/")
        assert alm.nombre == "Google Sheets"
    finally:
        storage.olvidar_almacenes()


# ============================================================================ el acceso (secretos [oauth])
def test_el_acceso_tambien_puede_venir_de_la_seccion_oauth(monkeypatch):
    import streamlit as st

    from utils import auth
    monkeypatch.setattr(st, "secrets", {"oauth": {"username": "equipo", "password": "clave-nueva"}})
    assert auth.verificar_credenciales("EQUIPO", "clave-nueva") and not auth.verificar_credenciales("Impacto", "glocal")
    monkeypatch.setattr(st, "secrets", {"auth": {"usuario": "otro", "clave": "x"}, "oauth": {"username": "equipo", "password": "y"}})
    assert auth.verificar_credenciales("otro", "x")                                    # [auth] tiene prioridad
    assert auth.verificar_clave_admin("glocal-admin")
