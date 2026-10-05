# -*- coding: utf-8 -*-
"""Almacenamiento + repositorio sobre una copia temporal del Excel real."""
import json
import re

import openpyxl
import pytest

from utils import repo
from utils import schema as S
from utils.repo import Cambio, ErrorOperacion
from utils.storage import ErrorAlmacen, LocalStorage

USUARIO = "Francis"


def _valor(almacen, ide, campo):
    base = almacen.leer_base_texto().set_index("id_evento", drop=False)
    return str(base.at[ide, campo])


def _sheet_names(almacen):
    wb = openpyxl.load_workbook(almacen.ruta, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


# ============================================================================ migración
def test_migracion_crea_columnas_hojas_y_no_toca_original(almacen, semilla):
    df = almacen.leer_base_texto()
    assert len(df) == 274
    assert list(df.columns[:len(S.COLUMNAS)]) == list(S.COLUMNAS)
    assert set(_sheet_names(almacen)) >= {S.HOJA_BASE, S.HOJA_LIBRO, *S.HOJAS_AUX}
    assert (df["origen"] == "historico").all()
    assert (df["fecha_carga"] != "").all()
    assert (df["id_evento"].str.match(r"^EV\d{4}$")).all()


def test_migracion_idempotente(almacen):
    firma = almacen.firma()
    almacen.asegurar()
    assert almacen.firma() == firma


def test_lectura_en_texto_canonico(almacen):
    df = almacen.leer_base_texto()
    assert set(df["Consultora"]) <= {"EIRL", "SpA", "Ltda", "0"}          # 0 (int) y '0' (str) unificados
    assert not df.isin(["nan", "NaT", "None"]).any().any()
    assert df["fecha_publicacion_web"].str.match(r"^\d{4}-\d{2}-\d{2}").all()
    assert df["num_enlaces_externos"].str.match(r"^\d+$").all()


def test_categorias_sembradas(almacen):
    cats = repo.categorias_df(almacen)
    assert set(cats["dimension"]) == set(S.DIMENSIONES_CATALOGO)
    macro = cats[cats["dimension"] == "categoria_macro"]["nombre"].tolist()
    assert "Territorio, medioambiente y sostenibilidad" in macro
    assert not any(n.lower() == "no aplica" for n in cats["nombre"])
    assert repo.catalogo_activo(almacen)["categoria_macro"]


def test_libro_de_codigos_generado(almacen):
    lc = almacen.leer_hoja_texto(S.HOJA_LIBRO)
    assert list(lc.columns) == ["columna", "descripcion", "tipo_variable", "opciones_respuesta", "fuente"]
    assert lc["columna"].tolist()[:49] == [c for c in S.COLUMNAS[:49]] or set(S.COLUMNAS) <= set(lc["columna"])
    assert set(S.COLUMNAS_NUEVAS) <= set(lc["columna"])
    fila = lc[lc["columna"] == "titulo"].iloc[0]
    assert "glocalminds.com" in fila["fuente"]                            # se conservó el texto escrito a mano
    # recuento calculado desde los datos (antes decía ~577)
    base = almacen.leer_base_texto()
    n = len({p for v in base["actores_normalizados"] for p in S.dividir_etiquetas(v)
             if p.lower() not in S.VALORES_NO_ETIQUETA})
    texto = lc[lc["columna"] == "actores_normalizados"].iloc[0]["opciones_respuesta"]
    assert str(n) in texto and "~577" not in texto
    # metodologías: se cuentan tras unificar los alias, igual que la vista de lectura de la app
    metodos = {S.ALIAS_METODOLOGIA.get(p.lower(), p) for v in base["metodologia"] for p in S.dividir_etiquetas(v)
               if p.lower() not in S.VALORES_NO_ETIQUETA}
    texto = lc[lc["columna"] == "metodologia"].iloc[0]["opciones_respuesta"]
    assert texto.startswith(f"{len(metodos)} metodologías") and len(metodos) == 29


# ============================================================================ transacciones
def test_transaccion_con_error_no_guarda(almacen):
    firma = almacen.firma()
    with pytest.raises(RuntimeError):
        with almacen.transaccion() as libro:
            libro.set_valor("EV0001", "titulo", "NO DEBE GUARDARSE")
            raise RuntimeError("falla")
    assert almacen.firma() == firma
    assert _valor(almacen, "EV0001", "titulo") != "NO DEBE GUARDARSE"
    assert not list(almacen.ruta.parent.glob("*.tmp"))


def test_archivo_ausente_da_error_claro(tmp_path):
    alm = LocalStorage(tmp_path / "no_existe.xlsx")
    with pytest.raises(ErrorAlmacen, match="No se encuentra"):
        alm.leer_base_texto()


def test_respaldos_se_podan_a_cinco(almacen, monkeypatch):
    import utils.storage as st
    contador = iter(range(100))

    class _Reloj:
        @staticmethod
        def now():
            from datetime import datetime, timedelta
            return datetime(2026, 1, 1) + timedelta(seconds=next(contador))
    monkeypatch.setattr(st, "datetime", _Reloj)
    for _ in range(8):
        assert almacen.respaldar("t")
    assert len(almacen.listar_respaldos()) == st.MAX_RESPALDOS


# ============================================================================ cambios + historial
def test_cambio_registra_historial_y_sella(almacen):
    antes = _valor(almacen, "EV0001", "titulo")
    res = repo.aplicar_cambios([Cambio("EV0001", "titulo", antes, "Título nuevo")], USUARIO, almacen=almacen)
    assert res.ok and len(res.aplicados) == 1
    assert _valor(almacen, "EV0001", "titulo") == "Título nuevo"
    assert _valor(almacen, "EV0001", "editado_por") == USUARIO
    assert _valor(almacen, "EV0001", "fecha_edicion") != ""
    h = repo.historial_df(almacen, "EV0001")
    assert len(h) == 1
    fila = h.iloc[0]
    assert (fila["usuario"], fila["campo"], fila["valor_antes"], fila["valor_despues"], fila["accion"]) == \
           (USUARIO, "titulo", antes, "Título nuevo", S.ACC_EDITAR)


def test_cambio_sin_diferencia_no_escribe(almacen):
    firma = almacen.firma()
    actual = _valor(almacen, "EV0001", "titulo")
    res = repo.aplicar_cambios([Cambio("EV0001", "titulo", actual, actual)], USUARIO, almacen=almacen)
    assert res.n_aplicados == 0 and almacen.firma() == firma
    assert repo.historial_df(almacen).empty


def test_conflicto_si_otra_persona_cambio_la_celda(almacen):
    antes = _valor(almacen, "EV0001", "lugar")
    repo.aplicar_cambios([Cambio("EV0001", "lugar", antes, "Otro lugar")], "Ana", almacen=almacen)
    res = repo.aplicar_cambios([Cambio("EV0001", "lugar", antes, "Mi lugar")], "Beto", almacen=almacen)
    assert not res.ok and "Otra persona" in res.omitidos[0]["motivo"]
    assert _valor(almacen, "EV0001", "lugar") == "Otro lugar"


def test_campos_no_editables_y_obligatorios_vacios(almacen):
    res = repo.aplicar_cambios([Cambio("EV0001", "id_evento", "EV0001", "EV9999"),
                                Cambio("EV0001", "titulo", _valor(almacen, "EV0001", "titulo"), "")],
                               USUARIO, almacen=almacen)
    assert len(res.omitidos) == 2 and res.n_aplicados == 0
    assert _valor(almacen, "EV0001", "id_evento") == "EV0001"


def test_falta_nombre_del_editor(almacen):
    with pytest.raises(ErrorOperacion, match="nombre"):
        repo.aplicar_cambios([], "  ", almacen=almacen)


def test_tipos_se_escriben_tipados(almacen):
    ide = "EV0002"
    repo.aplicar_cambios([Cambio(ide, "fecha_publicacion_web", _valor(almacen, ide, "fecha_publicacion_web"), "12/07/2026"),
                          Cambio(ide, "es_duplicado_secundario", "False", "sí")], USUARIO, almacen=almacen)
    wb = openpyxl.load_workbook(almacen.ruta, read_only=True)
    ws = wb[S.HOJA_BASE]
    cab = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    fila = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))          # EV0002 es la fila 3
    wb.close()
    assert fila[cab.index("fecha_publicacion_web")].year == 2026
    assert fila[cab.index("es_duplicado_secundario")] is True
    assert _valor(almacen, ide, "fecha_publicacion_web") == "2026-07-12"


def test_celdas_no_tocadas_conservan_su_tipo(almacen):
    """Editar una celda no debe convertir a texto el resto del libro (p. ej. Consultora = 0 entero)."""
    ide = next(i for i in almacen.leer_base_texto()["id_evento"])
    repo.aplicar_cambios([Cambio(ide, "autor", "", "Alguien")], USUARIO, almacen=almacen)
    wb = openpyxl.load_workbook(almacen.ruta, read_only=True)
    ws = wb[S.HOJA_BASE]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()
    cab = filas[0]
    tipos = {type(f[cab.index("Consultora")]).__name__ for f in filas[1:]}
    assert "int" in tipos                                                       # los 0 numéricos siguen siendo int
    assert type(filas[1][cab.index("fecha_publicacion_web")]).__name__ == "datetime"


# ============================================================================ revertir
def test_revertir_cambio(almacen):
    antes = _valor(almacen, "EV0003", "titulo")
    repo.aplicar_cambios([Cambio("EV0003", "titulo", antes, "Cambiado")], "Ana", almacen=almacen)
    idc = repo.historial_df(almacen, "EV0003").iloc[0]["id_cambio"]
    res = repo.revertir_cambio(idc, "Beto", almacen=almacen)
    assert res.ok and _valor(almacen, "EV0003", "titulo") == antes
    h = repo.historial_df(almacen, "EV0003")
    assert len(h) == 2                                                          # no se borra nada
    ult = h.iloc[0]
    assert (ult["accion"], ult["revierte_a"], ult["usuario"]) == (S.ACC_REVERTIR, idc, "Beto")


def test_revertir_cambio_con_conflicto_y_forzado(almacen):
    v0 = _valor(almacen, "EV0003", "lugar")
    repo.aplicar_cambios([Cambio("EV0003", "lugar", v0, "A")], "Ana", almacen=almacen)
    id_a = repo.historial_df(almacen, "EV0003").iloc[0]["id_cambio"]
    repo.aplicar_cambios([Cambio("EV0003", "lugar", "A", "B")], "Ana", almacen=almacen)
    res = repo.revertir_cambio(id_a, "Beto", almacen=almacen)
    assert not res.ok and _valor(almacen, "EV0003", "lugar") == "B"
    res = repo.revertir_cambio(id_a, "Beto", almacen=almacen, forzar=True)
    assert res.ok and _valor(almacen, "EV0003", "lugar") == v0


def test_restaurar_a_una_fecha(almacen, monkeypatch):
    ide = "EV0004"
    t = iter(["2026-10-01T10:00:00", "2026-10-02T10:00:00", "2026-10-03T10:00:00"])
    monkeypatch.setattr(repo, "ahora", lambda: next(t))
    t0 = _valor(almacen, ide, "titulo")
    l0 = _valor(almacen, ide, "lugar")
    repo.aplicar_cambios([Cambio(ide, "titulo", t0, "T1")], "Ana", almacen=almacen)           # día 1
    repo.aplicar_cambios([Cambio(ide, "lugar", l0, "L1")], "Ana", almacen=almacen)            # día 2
    repo.aplicar_cambios([Cambio(ide, "titulo", "T1", "T2")], "Ana", almacen=almacen)         # día 3
    plan = repo.plan_restauracion(ide, "2026-10-01 12:00:00", almacen)
    assert {p["campo"]: (p["actual"], p["objetivo"]) for p in plan} == {"lugar": ("L1", l0), "titulo": ("T2", "T1")}
    monkeypatch.setattr(repo, "ahora", lambda: "2026-10-04T10:00:00")
    res = repo.restaurar_noticia(ide, "2026-10-01 12:00:00", "Beto", almacen=almacen)
    assert res.ok and _valor(almacen, ide, "titulo") == "T1" and _valor(almacen, ide, "lugar") == l0


def test_restaurar_noticia_creada_despues_falla(almacen):
    res = repo.agregar_filas([{"titulo": "Nueva", "url_noticia": "https://x.cl/n", "fuente": "glocalminds.com"}],
                             USUARIO, almacen=almacen)
    with pytest.raises(ErrorOperacion, match="se creó después"):
        repo.plan_restauracion(res.ids_creados[0], "2000-01-01 00:00:00", almacen)


# ============================================================================ altas y lotes
def _nueva(i=1, **extra):
    return {"titulo": f"Noticia {i}", "url_noticia": f"https://glocalminds.com/catalogo/n{i}/",
            "fuente": "glocalminds.com", "slug": f"n{i}", "wp_id": str(9000 + i), **extra}


def test_agregar_filas(almacen):
    res = repo.agregar_filas([_nueva(1), _nueva(2)], USUARIO, accion=S.ACC_CARGA, origen="archivo", almacen=almacen)
    assert res.ids_creados == ["EV0275", "EV0276"] and res.ok
    fila = almacen.leer_base_texto().set_index("id_evento").loc["EV0275"]
    assert (fila["origen"], fila["cargado_por"], fila["wp_id"], fila["es_duplicado_secundario"]) == \
           ("archivo", USUARIO, "9001", "False")
    assert fila["fecha_carga"] != "" and fila["editado_por"] == ""
    h = repo.historial_df(almacen, "EV0275").iloc[0]
    assert h["campo"] == repo.REGISTRO and json.loads(h["valor_despues"])["titulo"] == "Noticia 1"


def test_agregar_omite_duplicados_seguros_e_invalidos(almacen):
    repo.agregar_filas([_nueva(1)], USUARIO, almacen=almacen)
    res = repo.agregar_filas([
        _nueva(1),                                                # mismo wp_id/url/slug
        {"titulo": "Sin url", "fuente": "glocalminds.com"},        # falta obligatorio
        _nueva(2),
    ], USUARIO, almacen=almacen)
    assert res.ids_creados == ["EV0276"]
    assert len(res.omitidos) == 2
    assert "Ya existe" in res.omitidos[0]["motivo"] and "obligatorio" in res.omitidos[1]["motivo"]


def test_agregar_misma_tanda_se_detecta_a_si_misma(almacen):
    res = repo.agregar_filas([_nueva(5), _nueva(5)], USUARIO, almacen=almacen)
    assert len(res.ids_creados) == 1 and len(res.omitidos) == 1


def test_titulo_repetido_no_se_omite(almacen):
    """Nivel 4 (mismo título) es solo aviso: puede ser otro evento."""
    titulo = _valor(almacen, "EV0001", "titulo")
    res = repo.agregar_filas([_nueva(7, titulo=titulo)], USUARIO, almacen=almacen)
    assert res.ids_creados == ["EV0275"]


def test_deshacer_lote_de_altas_y_rehacer(almacen):
    res = repo.agregar_filas([_nueva(1), _nueva(2)], USUARIO, accion=S.ACC_SINCRONIZAR, origen="scraping", almacen=almacen)
    lote = res.lote_id
    lotes = repo.lotes_df(almacen)
    assert lotes.iloc[0]["lote_id"] == lote and lotes.iloc[0]["n_noticias"] == 2 and not lotes.iloc[0]["revertido"]
    plan = repo.plan_lote(lote, almacen)
    assert [p["tipo"] for p in plan] == ["eliminar", "eliminar"] and not any(p["conflicto"] for p in plan)

    res2 = repo.revertir_lote(lote, "Ana", almacen=almacen)
    assert sorted(res2.ids_eliminados) == ["EV0275", "EV0276"]
    assert len(almacen.leer_base_texto()) == 274
    assert repo.lotes_df(almacen).set_index("lote_id").loc[lote, "revertido"]
    assert repo.plan_lote(lote, almacen) == []                                  # ya revertido

    # deshacer el deshacer vuelve a crear las noticias con sus mismos IDs y datos
    res3 = repo.revertir_lote(res2.lote_id, "Ana", almacen=almacen)
    assert sorted(res3.ids_creados) == ["EV0275", "EV0276"]
    fila = almacen.leer_base_texto().set_index("id_evento").loc["EV0276"]
    assert fila["titulo"] == "Noticia 2" and fila["wp_id"] == "9002"


def test_deshacer_lote_de_ediciones_con_conflicto(almacen):
    ids = ["EV0010", "EV0011"]
    cambios = [Cambio(i, "lugar", _valor(almacen, i, "lugar"), f"L-{i}") for i in ids]
    res = repo.aplicar_cambios(cambios, "Ana", accion=S.ACC_CATEGORIA, almacen=almacen)
    repo.aplicar_cambios([Cambio("EV0011", "lugar", "L-EV0011", "Otro")], "Beto", almacen=almacen)   # conflicto
    plan = {p["id_evento"]: p for p in repo.plan_lote(res.lote_id, almacen)}
    assert plan["EV0010"]["conflicto"] == "" and plan["EV0011"]["conflicto"]
    r = repo.revertir_lote(res.lote_id, "Ana", almacen=almacen)
    assert len(r.aplicados) == 1 and len(r.omitidos) == 1
    assert _valor(almacen, "EV0011", "lugar") == "Otro"


def test_deshacer_lote_no_retira_noticia_editada_despues(almacen):
    res = repo.agregar_filas([_nueva(1)], USUARIO, almacen=almacen)
    repo.aplicar_cambios([Cambio("EV0275", "autor", "", "X")], "Beto", almacen=almacen)
    r = repo.revertir_lote(res.lote_id, "Ana", almacen=almacen)
    assert r.ids_eliminados == [] and r.omitidos and len(almacen.leer_base_texto()) == 275


# ============================================================================ notas, categorías, duplicados, exportaciones
def test_notas(almacen):
    repo.agregar_nota("EV0001", "Revisar la clasificación", "Ana", almacen)
    repo.agregar_nota("EV0001", "Segunda nota", "Beto", almacen)
    df = repo.notas_df(almacen, "EV0001")
    assert df["texto"].tolist() == ["Segunda nota", "Revisar la clasificación"]
    assert df["usuario"].tolist() == ["Beto", "Ana"]
    with pytest.raises(ErrorOperacion):
        repo.agregar_nota("EV0001", "   ", "Ana", almacen)
    with pytest.raises(ErrorOperacion):
        repo.agregar_nota("EV9999", "x", "Ana", almacen)


def test_categorias_crear_renombrar_asignar(almacen):
    repo.crear_categoria("categorias", "Categoría de prueba", "desc", USUARIO, almacen)
    with pytest.raises(ErrorOperacion, match="ya existe"):
        repo.crear_categoria("categorias", "categoría DE prueba", "", USUARIO, almacen)
    with pytest.raises(ErrorOperacion):
        repo.crear_categoria("categorias", "con;punto y coma", "", USUARIO, almacen)

    res = repo.asignar_categoria("categorias", "Categoría de prueba", ["EV0001", "EV0002"], USUARIO, "agregar", almacen)
    assert len(res.aplicados) == 2
    assert "Categoría de prueba" in S.dividir_etiquetas(_valor(almacen, "EV0001", "categorias"))
    assert repo.uso_categorias(almacen)[("categorias", "Categoría de prueba")] == 2

    r2 = repo.renombrar_categoria("categorias", "Categoría de prueba", "Categoría renombrada", USUARIO, almacen)
    assert len(r2.aplicados) == 2
    assert "Categoría renombrada" in S.dividir_etiquetas(_valor(almacen, "EV0002", "categorias"))
    assert "Categoría de prueba" not in S.dividir_etiquetas(_valor(almacen, "EV0002", "categorias"))
    nombres = repo.categorias_df(almacen)["nombre"].tolist()
    assert "Categoría renombrada" in nombres and "Categoría de prueba" not in nombres
    assert all(a["id_cambio"] for a in r2.aplicados) and len({a["id_cambio"] for a in r2.aplicados}) == 2

    # asignar es deshacible por lote
    rq = repo.asignar_categoria("categorias", "Categoría renombrada", ["EV0001", "EV0002"], USUARIO, "quitar", almacen)
    assert "Categoría renombrada" not in _valor(almacen, "EV0001", "categorias")
    repo.revertir_lote(rq.lote_id, USUARIO, almacen)
    assert "Categoría renombrada" in _valor(almacen, "EV0001", "categorias")

    repo.actualizar_categoria("categorias", "Categoría renombrada", almacen, activa=False)
    assert "Categoría renombrada" not in repo.catalogo_activo(almacen)["categorias"]


def test_decision_de_duplicados(almacen):
    res = repo.registrar_decision_duplicado("EV0076", "EV0099", "duplicado", USUARIO, secundaria="EV0099", almacen=almacen)
    assert res.ok and _valor(almacen, "EV0099", "es_duplicado_secundario") == "True"
    repo.registrar_decision_duplicado("EV0094", "EV0122", "distintas", USUARIO, almacen=almacen)
    d = repo.decisiones_duplicados(almacen)
    assert d["decision"].tolist() == ["duplicado", "distintas"]
    with pytest.raises(ErrorOperacion):
        repo.registrar_decision_duplicado("EV0001", "EV0002", "duplicado", USUARIO, secundaria="EV0003", almacen=almacen)


def test_registro_de_exportaciones(almacen):
    ide = repo.registrar_exportacion(USUARIO, "Fondo X", "xlsx", "seleccion", ["EV0001", "EV0002"],
                                     ["titulo", "url_noticia"], "texto=agua", almacen)
    df = repo.exportaciones_df(almacen)
    assert ide == "X000001" and df.iloc[0]["ids"] == "EV0001,EV0002" and df.iloc[0]["contexto"] == "Fondo X"


def test_restablecer_desde_semilla(almacen):
    repo.aplicar_cambios([Cambio("EV0001", "autor", "", "X")], USUARIO, almacen=almacen)
    respaldo = almacen.restablecer_desde_semilla()
    assert respaldo and respaldo.exists()
    assert _valor(almacen, "EV0001", "autor") == ""
    assert repo.historial_df(almacen).empty
    assert re.search(r"BACKUP_\d{8}_\d{6}_antes_de_restablecer", respaldo.name)
