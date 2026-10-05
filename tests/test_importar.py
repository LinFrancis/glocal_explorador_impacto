# -*- coding: utf-8 -*-
import io

import openpyxl
import pandas as pd
import pytest

from utils import dedupe, importar
from utils import schema as S
from utils.importar import ErrorImportacion


def _csv(texto: str, cod="utf-8") -> bytes:
    return texto.encode(cod)


# ----------------------------------------------------------------------------- columnas
def test_mapear_columnas_por_clave_etiqueta_y_alias():
    mapa = importar.mapear_columnas(["Título *", "url_noticia", "Sitio", "Fecha", "Resumen", "Categoría macro",
                                     "Columna rara", "id_evento", "wp_id", "titulo"])
    assert mapa["Título *"] == "titulo"
    assert mapa["url_noticia"] == "url_noticia"
    assert mapa["Sitio"] == "fuente" and mapa["Fecha"] == "fecha_publicacion_web"
    assert mapa["Resumen"] == "descripcion_catalogo" and mapa["Categoría macro"] == "categoria_macro"
    assert "Columna rara" not in mapa
    assert "id_evento" not in mapa and "wp_id" not in mapa                  # los genera la plataforma
    assert "titulo" not in mapa                                              # ya usado por «Título *»


def test_mapear_columnas_del_excel_original(semilla):
    wb = openpyxl.load_workbook(semilla, read_only=True)
    cab = [c for c in next(wb["Base_Datos"].iter_rows(values_only=True)) if c]
    wb.close()
    mapa = importar.mapear_columnas(cab)
    assert len(mapa) == 48 and "id_evento" not in mapa                       # 49 menos id_evento


# ----------------------------------------------------------------------------- lectura
def test_leer_csv_con_punto_y_coma_y_latin1():
    datos = _csv("titulo;url_noticia;lugar\nTaller de agua;https://glocalminds.com/a/;Valparaíso\n", "latin-1")
    df = importar.leer_archivo(datos, "x.csv")
    assert df.shape == (1, 3) and df.iloc[0]["lugar"] == "Valparaíso"


def test_leer_excel_todo_texto_y_hojas(tmp_path):
    ruta = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["titulo", "num"])
    ws.append(["A", 3])
    wb.create_sheet("Otra").append(["x"])
    wb.save(ruta)
    datos = ruta.read_bytes()
    assert importar.hojas_de(datos, "t.xlsx") == ["Datos", "Otra"] and importar.hojas_de(datos, "t.csv") == []
    df = importar.leer_archivo(datos, "t.xlsx", "Datos")
    assert df.iloc[0]["num"] == "3" and df.dtypes.astype(str).eq("str").all() or df.iloc[0]["titulo"] == "A"


@pytest.mark.parametrize("datos,nombre,texto", [
    pytest.param(b"titulo\n", "vacio.csv", "no tiene filas", id="vacio"),
    pytest.param(b"x" * (importar.MAX_BYTES + 1), "grande.csv", "pesa más", id="demasiado-grande"),
    pytest.param(b"a", "algo.txt", "Formato no admitido", id="formato"),
    pytest.param(b"no soy un excel", "roto.xlsx", "No se pudo", id="excel-roto"),
])
def test_errores_de_lectura_en_espanol(datos, nombre, texto):
    with pytest.raises(ErrorImportacion, match=texto):
        importar.leer_archivo(datos, nombre)


def test_demasiadas_filas():
    cuerpo = "titulo\n" + "\n".join(f"T{i}" for i in range(importar.MAX_FILAS + 1))
    with pytest.raises(ErrorImportacion, match="máximo"):
        importar.leer_archivo(_csv(cuerpo), "m.csv")


# ----------------------------------------------------------------------------- filas y evaluación
def test_fuente_se_infiere_del_dominio():
    assert importar.inferir_fuente("https://www.glocalminds.com/catalogo/x/") == "glocalminds.com"
    assert importar.inferir_fuente("fundacionglocal.org/mingamar") == "fundacionglocal.org"
    assert importar.inferir_fuente("https://otro.cl/x") == ""
    df = pd.DataFrame({"Título": ["A"], "Enlace": ["https://glocalminds.com/a/"]})
    filas = importar.filas_desde_df(df, importar.mapear_columnas(list(df.columns)))
    assert filas == [{"titulo": "A", "url_noticia": "https://glocalminds.com/a/", "fuente": "glocalminds.com"}]


def test_evaluar_estados(almacen):
    base = almacen.leer_base_texto().to_dict("records")
    existente = base[0]
    filas = [
        {"titulo": "Nueva y limpia", "url_noticia": "https://glocalminds.com/n1/", "fuente": "glocalminds.com",
         "fecha_publicacion_web": "2026-09-15"},
        {"titulo": "Con aviso", "url_noticia": "https://glocalminds.com/n2/", "fuente": "glocalminds.com",
         "fecha_publicacion_web": "ayer por la tarde"},
        {"titulo": "", "url_noticia": "https://glocalminds.com/n3/", "fuente": "glocalminds.com"},
        {"titulo": "Cualquiera", "url_noticia": existente["url_noticia"], "fuente": existente["fuente"]},
        {"titulo": existente["titulo"], "url_noticia": "https://glocalminds.com/otra-url/", "fuente": "glocalminds.com"},
        {"titulo": "Nueva y limpia", "url_noticia": "https://glocalminds.com/n1/", "fuente": "glocalminds.com",
         "fecha_publicacion_web": "2026-09-15"},
    ]
    ev = importar.evaluar(filas, dedupe.Indice(base))
    assert [e.estado for e in ev] == [
        importar.ESTADO_NUEVA, importar.ESTADO_AVISOS, importar.ESTADO_ERROR,
        importar.ESTADO_DUPLICADA, importar.ESTADO_SOSPECHOSA, importar.ESTADO_DUPLICADA,
    ]
    assert [e.importable for e in ev] == [True, True, False, False, False, False]
    assert "en este archivo" in ev[5].detalle and "obligatorio" in ev[2].detalle
    assert ev[3].detalle and existente["id_evento"] in ev[3].detalle


def test_plantilla_se_puede_leer_de_vuelta():
    datos = importar.plantilla_xlsx()
    assert importar.hojas_de(datos, "p.xlsx") == ["Noticias", "Instrucciones"]
    df = importar.leer_archivo(datos, "p.xlsx", "Noticias")
    mapa = importar.mapear_columnas(list(df.columns))
    assert set(mapa.values()) == set(importar.PLANTILLA_COLUMNAS)
    filas = importar.filas_desde_df(df, mapa)
    assert len(filas) == 1 and filas[0]["titulo"].startswith("Taller de gobernanza")
    assert importar.evaluar(filas, dedupe.Indice())[0].importable
    ins = openpyxl.load_workbook(io.BytesIO(datos))["Instrucciones"]
    assert ins["A2"].value == S.etiqueta("titulo") and ins["B2"].value == "Sí"
