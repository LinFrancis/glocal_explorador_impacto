# -*- coding: utf-8 -*-
import io

import openpyxl
import pandas as pd
import pytest
from docx import Document

from tests.apptest_util import apuntar_a
from utils import data, export
from utils import schema as S


@pytest.fixture
def df(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    return data.load_noticias()


def _hojas(datos: bytes) -> dict[str, list[list]]:
    wb = openpyxl.load_workbook(io.BytesIO(datos))
    return {ws.title: [[c.value for c in fila] for fila in ws.iter_rows()] for ws in wb.worksheets}


def _texto_word(datos: bytes) -> str:
    doc = Document(io.BytesIO(datos))
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        partes += [c.text for fila in t.rows for c in fila.cells]
    return "\n".join(partes)


# ----------------------------------------------------------------------------- catálogo
def test_catalogo_exportable_agrupado_y_con_derivados(df):
    cat = export.catalogo_exportable(df.columns)
    claves = [k for k, _, _ in cat]
    assert {"titulo", "anio", "pais", "categoria_macro", "url_noticia"} <= set(claves)
    assert not set(claves) & set(S.COLUMNAS_INTERNAS) - {"anio", "pais"}
    assert [g for _, _, g in cat] == sorted((g for _, _, g in cat), key=S.GRUPOS.index)       # agrupado en orden
    assert export.etiqueta("anio") == "Año" and export.etiqueta("categoria_macro") == "Categoría macro"
    assert export.columnas_validas(["titulo", "no_existe", "anio"], df.columns) == ["titulo", "anio"]


# ----------------------------------------------------------------------------- Excel / CSV
def test_excel_completo_tiene_resumen_y_datos_sin_columnas_internas(df):
    hojas = _hojas(export.experiences_to_excel(df.head(3)))
    assert list(hojas) == ["Resumen", "Datos completos"]
    cab = hojas["Resumen"][0]
    assert cab[:3] == ["Título", "Año", "País"] and "Enlace" in cab and len(hojas["Resumen"]) == 4
    completos = hojas["Datos completos"][0]
    for interna in ("_palabras_busqueda", "item", "completitud", "enlaces_externos_lista", "fecha_parsed", "Latitudes_x"):
        assert interna not in completos
    assert "Título" in completos and "Latitudes" in completos and not any(str(c).endswith("_primary") for c in completos)
    assert "ID" in completos and "ID en WordPress" in completos


def test_excel_con_columnas_elegidas_respeta_orden_y_hoja_unica(df):
    cols = ["url_noticia", "titulo", "categoria_macro", "anio"]
    hojas = _hojas(export.experiences_to_excel(df.head(2), cols))
    assert list(hojas) == ["Noticias"]
    assert hojas["Noticias"][0] == ["Enlace", "Título", "Categoría macro", "Año"]
    fila = hojas["Noticias"][1]
    assert fila[1] == df.iloc[0]["titulo"] and str(fila[3]).isdigit()
    assert ";" not in (fila[2] or "") or "; " in fila[2]                                      # etiquetas limpias


def test_valores_exportados_son_texto_legible(df):
    hojas = _hojas(export.experiences_to_excel(df.head(2), ["fecha_publicacion_web", "es_duplicado_secundario", "num_enlaces_externos"]))
    f = hojas["Noticias"][1]
    assert f[0][:4].isdigit() and "00:00:00" not in f[0] and f[1] in ("True", "False") and f[2].isdigit()


def test_csv(df):
    texto = export.experiences_to_csv(df.head(2), ["titulo", "url_noticia"]).decode("utf-8-sig")
    lineas = texto.strip().splitlines()
    assert lineas[0] == "Título,Enlace" and len(lineas) == 3
    completo = pd.read_csv(io.StringIO(export.experiences_to_csv(df.head(2)).decode("utf-8-sig")))
    assert "_palabras_busqueda" not in completo.columns and completo.shape[0] == 2


def test_caracteres_de_control_no_rompen_la_exportacion(df):
    sucio = df.head(1).copy()
    sucio["titulo"] = "Título\x0b con\x00 caracteres\x1f raros"
    sucio["contenido_completo"] = "Texto\x0cdel contenido"
    assert _hojas(export.experiences_to_excel(sucio, ["titulo"]))["Noticias"][1][0] == "Título con caracteres raros"
    assert "Título con caracteres raros" in _texto_word(export.experiences_to_word(sucio))


# ----------------------------------------------------------------------------- Word
def test_word_completo(df):
    texto = _texto_word(export.experiences_to_word(df.head(2), contexto="Fondo X"))
    assert "Contexto de la selección: Fondo X" in texto and "2 experiencia(s)" in texto
    assert df.iloc[0]["titulo"] in texto and "Categoría macro:" in texto and "Enlace:" in texto
    assert "Resumen" in texto                                                                  # tabla resumen


def test_word_solo_campos_elegidos(df):
    datos = export.experiences_to_word(df.head(2), campos=["titulo", "url_noticia"], incluir_tabla=False)
    texto = _texto_word(datos)
    assert "Enlace:" in texto and df.iloc[0]["titulo"] in texto
    assert "Categoría macro:" not in texto and "Texto completo" not in texto and "Año:" not in texto
    solo_cat = _texto_word(export.experiences_to_word(df.head(1), campos=["categoria_macro", "anio"], incluir_tabla=False))
    assert "Categoría macro:" in solo_cat and "Año:" in solo_cat and "Enlace:" not in solo_cat


def test_word_con_logos_en_portada_y_encabezado(df):
    sin = Document(io.BytesIO(export.experiences_to_word(df.head(1), con_logos=False)))
    con = Document(io.BytesIO(export.experiences_to_word(df.head(1), con_logos=True)))
    assert len(sin.inline_shapes) == 0
    assert len(con.inline_shapes) == 2                                                        # portada: ambos logos
    header = con.sections[0].header
    assert len(header.part.element.xpath(".//w:drawing")) == 2                                # encabezado: ambos logos
    assert con.sections[0].different_first_page_header_footer


def test_word_sin_archivos_de_logo_no_falla(df, monkeypatch, tmp_path):
    monkeypatch.setattr(export, "LOGO_GLOCALMINDS_PNG", tmp_path / "no.png")
    monkeypatch.setattr(export, "LOGO_FUNDACION_PNG", tmp_path / "no2.png")
    assert len(Document(io.BytesIO(export.experiences_to_word(df.head(1), con_logos=True))).inline_shapes) == 0
