# -*- coding: utf-8 -*-
"""Equipo (nombres registrados), tabla validada por el libro de códigos, ayuda de Inicio, marco teórico con
variables propias, sincronización sin modo de prueba y recarga del código."""
import subprocess
import sys

import pandas as pd
import pytest

from tests.apptest_util import RAIZ, apuntar_a, iniciar_sesion, nueva_app
from utils import edicion, repo
from utils import schema as S
from utils.repo import Cambio, ErrorOperacion


def _boton(at, etiqueta, indice=0):
    return [b for b in at.button if b.label == etiqueta][indice]


def _tablas(at, columna):
    return [t.value for t in at.dataframe if columna in getattr(t.value, "columns", [])]


def _fila(almacen, ide):
    return almacen.leer_base_texto().set_index("id_evento", drop=False).loc[ide]


# ============================================================================ equipo: lógica
def test_registrar_editor_unifica_grafias(almacen):
    assert repo.editores(almacen) == []
    assert repo.registrar_editor("  Ana   Soto ", almacen) == "Ana Soto"
    assert repo.registrar_editor("ana soto", almacen) == "Ana Soto"               # misma persona: no se duplica
    assert repo.registrar_editor("María José", almacen) == "María José"
    assert repo.registrar_editor("maria jose", almacen) == "María José"           # sin tildes ni mayúsculas
    assert repo.editores(almacen) == ["Ana Soto", "María José"]
    hoja = almacen.leer_hoja_texto(S.HOJA_EDITORES)
    assert hoja["nombre"].tolist() == ["Ana Soto", "María José"] and hoja["fecha_hora"].ne("").all()


@pytest.mark.parametrize("nombre", ["", "   ", "A", None])
def test_registrar_editor_rechaza_nombres_vacios(almacen, nombre):
    with pytest.raises(ErrorOperacion, match="al menos 2 letras"):
        repo.registrar_editor(nombre, almacen)
    assert repo.editores(almacen) == []


def test_editores_incluye_a_quienes_ya_figuran_en_el_historial(almacen):
    repo.aplicar_cambios([Cambio("EV0001", "autor", "", "X")], "Luis Pérez", almacen=almacen)
    repo.registrar_editor("Ana Soto", almacen)
    assert repo.editores(almacen) == ["Ana Soto", "Luis Pérez"]
    assert repo.registrar_editor("luis perez", almacen) == "Luis Pérez"            # reconoce al del historial


def test_quitar_editor_no_toca_el_historial(almacen):
    repo.registrar_editor("Ana Sotto", almacen)                                    # mal escrito
    repo.aplicar_cambios([Cambio("EV0001", "autor", "", "X")], "Ana Sotto", almacen=almacen)
    msg = repo.quitar_editor("ana sotto", almacen)
    assert "se quitó" in msg and almacen.leer_hoja_texto(S.HOJA_EDITORES).empty
    assert repo.historial_df(almacen).iloc[0]["usuario"] == "Ana Sotto"            # el pasado queda como estaba
    with pytest.raises(ErrorOperacion, match="no está en la lista"):
        repo.quitar_editor("Nadie", almacen)


def test_equipo_df(almacen):
    repo.registrar_editor("Ana Soto", almacen)
    repo.aplicar_cambios([Cambio("EV0001", "autor", "", "X"), Cambio("EV0002", "autor", "", "Y")], "Ana Soto", almacen=almacen)
    df = repo.equipo_df(almacen)
    assert list(df.columns) == ["Nombre", "Registrado", "Último cambio", "Cambios"]
    assert df.iloc[0]["Nombre"] == "Ana Soto" and df.iloc[0]["Cambios"] == 2 and df.iloc[0]["Último cambio"]


def test_la_hoja_editores_se_crea_en_libros_antiguos(almacen):
    import openpyxl
    wb = openpyxl.load_workbook(almacen.ruta)
    del wb[S.HOJA_EDITORES]
    wb.save(almacen.ruta)
    assert S.HOJA_EDITORES not in openpyxl.load_workbook(almacen.ruta).sheetnames
    assert repo.registrar_editor("Ana Soto", almacen) == "Ana Soto"                # la hoja se crea sola al escribir
    assert repo.editores(almacen) == ["Ana Soto"]


# ============================================================================ tabla: libro de códigos
@pytest.fixture
def datos(almacen, monkeypatch):
    apuntar_a(monkeypatch, almacen)
    from utils.data import libro_de_codigos, opciones_de_campos
    base = almacen.leer_base_texto()
    return base, opciones_de_campos(), libro_de_codigos()


def test_opciones_incluyen_los_codigos_especiales_del_libro(datos):
    base, opciones, libro = datos
    ops = edicion.opciones_para_tabla(opciones, base, libro)
    assert "No aplica" in ops["eje_gcaa"] and "No aplica" in ops["atributos_resiliencia"]
    assert "No especificado" in ops["metodologia"]
    assert "No aplica" not in ops["categoria_macro"] and "No aplica" not in ops["categorias"]       # el libro no los admite ahí
    assert set(ops["tipo_informacion"]) >= {"Evidencia de acción realizada", "Testimonio o entrevista", "Descripción de programa"}
    assert "lugar" not in ops and "titulo" not in ops                                               # texto libre: sin lista


def test_tipo_informacion_toma_las_opciones_oficiales_del_libro(datos):
    _, opciones, libro = datos
    assert len(opciones["tipo_informacion"]) == 6
    assert "Reflexión u opinión" in opciones["tipo_informacion"]


def test_preparar_df_y_cambios_con_listas(datos):
    base, opciones, libro = datos
    ops = edicion.opciones_para_tabla(opciones, base, libro)
    df = edicion.preparar_df(base, ops)
    assert isinstance(df.iloc[0]["categoria_macro"], list) and df.iloc[0]["categoria_macro"] == S.dividir_etiquetas(base.iloc[0]["categoria_macro"])
    # la misma lista (aunque llegue con otro formato) no es un cambio; una lista distinta sí
    igual = S.dividir_etiquetas(base.iloc[0]["categoria_macro"])
    assert edicion.cambios_desde_edicion(base, {0: {"categoria_macro": igual}}) == []
    c = edicion.cambios_desde_edicion(base, {0: {"categoria_macro": igual + ["Nueva"]}})
    assert len(c) == 1 and c[0].despues.endswith("; Nueva") and c[0].antes == base.iloc[0]["categoria_macro"]
    assert edicion.texto_de_celda(["a", "b", "a"]) == "a; b" and edicion.texto_de_celda([]) == ""


def test_column_config_usa_selectores_y_el_libro_en_la_ayuda(datos):
    base, opciones, libro = datos
    ops = edicion.opciones_para_tabla(opciones, base, libro)
    cfg = edicion.column_config(ops, libro)
    assert set(S.COLUMNAS) | {edicion.COL_COMPLETITUD} == set(cfg)
    assert cfg["categoria_macro"]["type_config"]["type"] == "multiselect"
    assert cfg["categoria_macro"]["type_config"].get("accept_new_options") in (None, False)       # lista cerrada
    assert cfg["metodologia"]["type_config"].get("accept_new_options") is True                    # vocabulario abierto
    assert cfg["tipo_informacion"]["type_config"]["type"] == "selectbox"
    assert cfg["enfoque_genero"]["type_config"].get("validate")                                    # formato «No» / «Sí: …»
    assert "Libro de códigos" in cfg["eje_gcaa"]["help"] and "No aplica" in cfg["eje_gcaa"]["help"]
    assert cfg["titulo"]["type_config"]["type"] == "text"                                          # el texto libre sigue siendo texto


def test_errores_de_codigos(datos):
    base, opciones, libro = datos
    ops = edicion.opciones_para_tabla(opciones, base, libro)
    macro = base.iloc[0]["categoria_macro"]
    cambios = [
        Cambio("EV0001", "categoria_macro", macro, macro + "; Inventada"),                      # lista cerrada: etiqueta que no existe
        Cambio("EV0001", "categoria_macro", macro, macro),                                       # sin cambio real (no debería llegar)
        Cambio("EV0002", "metodologia", base.iloc[1]["metodologia"], "Método nuevo de verdad"),  # vocabulario abierto: se admite
        Cambio("EV0003", "tipo_informacion", "Evidencia de acción realizada", "Cosa rara"),      # opción única fuera de lista
        Cambio("EV0004", "enfoque_genero", "No", "tal vez"),                                      # formato del libro
        Cambio("EV0005", "enfoque_genero", "No", "Sí: mujeres rurales"),                          # formato correcto
        Cambio("EV0006", "eje_gcaa", "A1", "No aplica"),                                          # código especial admitido
    ]
    errores = edicion.errores_de_codigos(cambios, ops)
    assert set(errores) == {("EV0001", "categoria_macro"), ("EV0003", "tipo_informacion"), ("EV0004", "enfoque_genero")}
    assert "«Inventada»" in errores[("EV0001", "categoria_macro")]
    assert "«No» o «Sí:" in errores[("EV0004", "enfoque_genero")]


def test_los_valores_que_ya_tenia_la_celda_no_se_cuestionan(datos):
    base, opciones, libro = datos
    ops = edicion.opciones_para_tabla(opciones, base, libro)
    # una celda con una etiqueta antigua fuera de la lista puede cambiar otra etiqueta sin que se reclame por la vieja
    antes = "Etiqueta antigua; " + S.dividir_etiquetas(base.iloc[0]["categoria_macro"])[0]
    nuevo = antes + "; " + S.dividir_etiquetas(base.iloc[1]["categoria_macro"])[0]
    assert edicion.errores_de_codigos([Cambio("EV0001", "categoria_macro", antes, nuevo)], ops) == {}


def test_variables_propias_son_listas_cerradas_en_la_tabla(almacen, monkeypatch):
    apuntar_a(monkeypatch, almacen)
    v = repo.crear_variable("Tamaño del proyecto", "opcion", ["Pequeño", "Mediano", "Grande"], "", "Ana", almacen)
    m = repo.crear_variable("Enfoques", "etiquetas", ["Agua", "Energía"], "", "Ana", almacen)
    from utils.data import libro_de_codigos, opciones_de_campos
    base = almacen.leer_base_texto()
    ops = edicion.opciones_para_tabla(opciones_de_campos(), base, libro_de_codigos())
    assert ops[v.clave] == ["Pequeño", "Mediano", "Grande"] and ops[m.clave] == ["Agua", "Energía"]
    cfg = edicion.column_config(ops, libro_de_codigos())
    assert cfg[v.clave]["type_config"]["type"] == "selectbox" and cfg[m.clave]["type_config"]["type"] == "multiselect"
    err = edicion.errores_de_codigos([Cambio("EV0001", v.clave, "", "Gigante"), Cambio("EV0001", m.clave, "", "Agua; Fuego"),
                                      Cambio("EV0002", v.clave, "", "Mediano")], ops)
    assert set(err) == {("EV0001", v.clave), ("EV0001", m.clave)}
    # y el libro de códigos ya trae la variable como una columna más, con sus opciones
    fila = libro_de_codigos().set_index("columna").loc[v.clave]
    assert fila["opciones_respuesta"] == "Pequeño; Mediano; Grande" and fila["tipo_variable"] == "Categórica única"


# ============================================================================ tabla: página
@pytest.fixture
def app(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    at.switch_page("app_pages/tabla.py").run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def _editar(at, filas):
    at.session_state["tabla_editor_0"] = {"edited_rows": filas, "added_rows": [], "deleted_rows": []}
    at.run()
    assert not at.exception, [e.value for e in at.exception]


def test_tabla_guarda_una_categoria_elegida_de_la_lista(app, almacen):
    macro = _fila(almacen, "EV0001")["categoria_macro"]
    otra = next(o for o in repo.catalogo_activo(almacen)["categoria_macro"] if o not in S.dividir_etiquetas(macro))
    edicion_grilla = {0: {"categoria_macro": S.dividir_etiquetas(macro) + [otra]}}
    _editar(app, edicion_grilla)
    assert not app.error
    app.session_state["tabla_editor_0"] = {"edited_rows": edicion_grilla, "added_rows": [], "deleted_rows": []}
    _boton(app, "Guardar 1 cambio(s)").click()
    app.run()
    assert not app.exception, [e.value for e in app.exception]
    assert _fila(almacen, "EV0001")["categoria_macro"] == macro + "; " + otra
    h = repo.historial_df(almacen).iloc[0]
    assert (h["campo"], h["valor_antes"], h["usuario"]) == ("categoria_macro", macro, "Francis")


def test_tabla_no_deja_guardar_lo_que_rompe_el_libro_de_codigos(app, almacen):
    _editar(app, {0: {"categoria_macro": ["Inventada"]}, 1: {"enfoque_genero": "tal vez"}})
    assert any("no respetan el libro de códigos" in e.value for e in app.error)
    assert any("«Inventada»" in c.value for c in app.caption) and any("«No» o «Sí:" in c.value for c in app.caption)
    guardar = next(b for b in app.button if b.label.startswith("Guardar"))
    assert guardar.disabled
    assert _fila(almacen, "EV0001")["categoria_macro"] != "Inventada" and repo.historial_df(almacen).empty
    # descartar deja la grilla limpia y se puede seguir
    _boton(app, "Descartar cambios").click()
    app.run()
    assert not app.error and any("Sin cambios pendientes" in c.value for c in app.caption)


def test_tabla_explica_que_las_categorias_se_eligen(app):
    assert any("se eligen de su lista" in c.value for c in app.caption)


# ============================================================================ Inicio: ayuda
def test_inicio_tiene_la_ayuda_cerrada_con_todo_lo_que_se_puede_hacer(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    assert not at.exception, [e.value for e in at.exception]
    ayuda = next((e for e in at.expander if "Ayuda: qué puedes hacer" in e.label), None)
    assert ayuda is not None, [e.label for e in at.expander]
    assert ayuda.proto.expanded is False                                        # cerrada: no estorba a quien ya conoce la plataforma
    assert [t.label for t in ayuda.tabs] == ["Primeros pasos", "Explorar y exportar", "Editar información", "Cargar y sincronizar",
                                             "Administrar y referencia", "Claves y datos"]
    texto = "\n".join(m.value for t in ayuda.tabs for m in t.markdown)
    for modulo in ("Explorador Glocal", "Fichas de noticia", "Tabla de datos", "Cargar información", "Sincronizar con la web",
                   "Historial de cambios", "Administración", "Marco teórico", "Glosario", "variables propias", "libro de códigos",
                   "carpeta del proyecto", "Editando como", "clave de administración", "Excel, Word o CSV"):
        assert modulo.lower() in texto.lower(), modulo
    assert "modo de prueba" not in texto.lower()


# ============================================================================ marco teórico y libro de códigos
def test_marco_teorico_registra_las_variables_propias(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    v = repo.crear_variable("Tamaño del proyecto", "opcion", ["Pequeño", "Mediano", "Grande"],
                            "Según el número de personas participantes", "Ana", almacen)
    repo.asignar_categoria(v.clave, "Grande", ["EV0001", "EV0002"], "Ana", "agregar", almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    at.switch_page("app_pages/marco_teorico.py").run()
    assert not at.exception, [e.value for e in at.exception]
    md = "\n".join(m.value for m in at.markdown)
    assert "Variables propias del equipo" in md and "Tamaño del proyecto" in md
    assert "Según el número de personas participantes" in md
    assert "Pequeño (0) · Mediano (0) · Grande (2)" in md                         # opciones de respuesta con su uso
    assert any(c.value == v.clave for c in at.code)
    libro = _tablas(at, "Opciones de respuesta")[0]
    fila = libro[libro["Columna"] == v.clave].iloc[0]                             # el libro de códigos la trae como una columna más
    assert fila["Variable"] == "Tamaño del proyecto" and fila["Opciones de respuesta"] == "Pequeño; Mediano; Grande"
    assert len(libro) == len(S.COLUMNAS)


def test_marco_teorico_sin_variables_invita_a_crearlas(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    at.switch_page("app_pages/marco_teorico.py").run()
    assert not at.exception, [e.value for e in at.exception]
    assert any("Todavía no hay variables propias" in i.value for i in at.info)


# ============================================================================ administración: equipo
@pytest.fixture
def admin(monkeypatch, almacen):
    apuntar_a(monkeypatch, almacen)
    at = nueva_app()
    iniciar_sesion(at, nombre="Francis")
    at.switch_page("app_pages/administracion.py").run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def test_admin_tiene_la_pestana_equipo_con_los_nombres(admin):
    assert len(admin.tabs) == 9 and admin.tabs[-1].label.endswith("Equipo")
    equipo = _tablas(admin, "Último cambio")[0]
    assert equipo["Nombre"].tolist() == ["Francis"]                                  # quien inició sesión quedó registrado


def test_admin_agrega_un_nombre_al_equipo(admin, almacen):
    admin.text_input(key="eq_nuevo").set_value("Ana Soto")
    [b for b in admin.button if b.label == "Agregar"][-1].click()                   # el de la pestaña Equipo es el último
    admin.run()
    assert not admin.exception, [e.value for e in admin.exception]
    assert repo.editores(almacen) == ["Ana Soto", "Francis"]
    assert any("se agregó al equipo" in s.value for s in admin.success)


def test_admin_quitar_nombre_pide_la_clave(admin, almacen):
    repo.registrar_editor("Ana Sotto", almacen)
    admin.run()
    admin.selectbox(key="eq_quitar").select("Ana Sotto").run()
    boton = next(b for b in admin.button if b.label == "Quitar «Ana Sotto»")
    boton.click()
    admin.run()
    assert not admin.exception, [e.value for e in admin.exception]
    assert "Ana Sotto" in repo.editores(almacen)                                    # abrir el diálogo no quita nada todavía


# ============================================================================ recarga del código
def test_el_servidor_descarta_modulos_viejos_al_cambiar_el_codigo():
    """Un servidor que ya tenía módulos viejos en memoria (Streamlit no siempre ve los cambios en Google Drive) los
    descarta y vuelve a importarlos, en vez de fallar con «no tiene el atributo…»."""
    codigo = (
        "import os, sys\n"
        "os.environ.pop('GLOCAL_RECARGA_CODIGO', None)\n"
        f"sys.path.insert(0, r'{RAIZ}')\n"
        "import utils.schema as viejo\n"
        "viejo.MARCA_VIEJA = True\n"
        "from streamlit.testing.v1 import AppTest\n"
        f"at = AppTest.from_file(r'{RAIZ / 'Inicio.py'}', default_timeout=90).run()\n"
        "nuevo = sys.modules['utils.schema']\n"
        "assert not at.exception, [e.value for e in at.exception]\n"
        "assert nuevo is not viejo and not hasattr(nuevo, 'MARCA_VIEJA'), 'no se recargó el módulo'\n"
        "assert hasattr(nuevo, 'NUMERO') and nuevo.CAMPO['carpeta_proyecto']\n"
        "print('RECARGA_OK')\n"
    )
    r = subprocess.run([sys.executable, "-c", codigo], capture_output=True, text=True, timeout=240)
    assert "RECARGA_OK" in r.stdout, r.stderr[-1500:]
